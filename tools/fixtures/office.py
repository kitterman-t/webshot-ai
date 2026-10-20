#!/usr/bin/env python3
"""Regenerate the office-format fixtures under `tests/fixtures/office/`.

    uv run --extra office python tools/fixtures/office.py

A committed `.docx` is a zip full of XML: unreviewable in a diff, and
unmaintainable once whoever made it has moved on.  This script is the fixture —
the files it writes are build output that happens to be committed so the test
suite does not need the office extra installed to *exist*, only to run.

Each fixture carries a sentinel string that occurs nowhere else in
the repository, so `tests/test_office_formats.py` can prove the text reached
the bundle through the converter rather than through a filename.  The one
exception is `equation_archive.zip`, which holds the equation deck and report
as they are, and so the deck's sentinel.

DOCX and EPUB are written by hand, because the packages that read them
(mammoth, MarkItDown's EPUB converter) cannot write them and pulling in a
writer for each would add dependencies no shipped code path uses.  Both
formats are small, declared containers; the XML below is the minimum each
specification requires.  PPTX and XLSX use python-pptx and openpyxl, which
`webshot[office]` already installs to read them.

The two `equation_*` fixtures hold what those readers pass over (docs/09 P27):
an equation, `mc:AlternateContent`, SmartArt, an embedded object, and an
equation MarkItDown's LaTeX converter refuses.  Every word in them is made up.
`equation_archive.zip` holds both, to show the same rewrite reaching a zip's
members (docs/09 P29).

The five `markup_*` fixtures hold text that HTML or Markdown would read as
structure — an `<img>` aimed at a `.invalid` host, `# heading`, `3.`, `===`,
`: definition`, a link, alt text and a link target that close their own
syntax — in every place each format's reader takes text from (docs/09 P28).

`left_out_archive.zip` holds one member of each kind a zip's reader leaves
out of the page, each of which the manifest must name (docs/09 P33).

`adjacent_lists_book.epub` holds lists with nothing between them but markup
that writes no text, which Python-Markdown read as one list (docs/09 P32).

`late_utf8.csv` and `late_utf8_archive.zip` are UTF-8 text whose first
character outside ASCII comes after the 64 KB MarkItDown guesses an encoding
from, so that guess is ASCII (docs/09 P35).  Their sentinels are that
character's words.

The three `nested_*` fixtures hold lists nested in lists, at every depth and
marker width a reader writes, which Python-Markdown read at their parent's
level (docs/09 P40); `nested_archive.zip` holds an HTML page and the Word file.

The three `code_items_*` fixtures hold a `<pre>` in list items at every depth
and next to everything else an item holds, whose lines Python-Markdown ran
together into one (docs/09 P44).  The Word file carries a mammoth style map,
the one way a `.docx` reaches markdownify with a `<pre>`.

The two `definition_*` fixtures hold definition lists in list items, each
definition holding more than its first paragraph, which Python-Markdown read
into the item instead (docs/09 P62); `definition_archive.zip` holds an HTML
page and the book.

The three `inline_code_*` fixtures hold a `<pre>` in a link, emphasis, struck
text, a subscript, a quotation and inline code, and inline code over several
lines: markdownify wrote each into a line of text, where Python-Markdown read
the code as Markdown, and an `<img>` aimed at a `.invalid` host typed in it
became an element (docs/09 P59).

The three `quoted_code_*` fixtures hold a `<pre>` in quotations and in
definitions, nested in each other and next to what else they hold, whose
lines Python-Markdown ran together the same way, and whose markup after a
blank line it read as HTML (docs/09 P47).

The three `edge_space_*` fixtures — Word, an EPUB, and an HTML page in a zip —
hold paragraphs whose words are kept apart only by a Unicode space at the edge
of bold, italic, a link, code and the rest (docs/09 P43).

The three `whitespace_run_*` fixtures hold paragraphs whose words are kept
apart only by ASCII whitespace or a line break in formatted text: a run of it
alone, a newline or a `<br>` at the run's edge, and a space at the edge of a
link that MarkItDown writes without its markup (docs/09 P49).

The three `adjacent_blocks_*` fixtures hold quotations and definition lists
side by side, which Python-Markdown read as one (docs/09 P63): an EPUB of two
chapters, a Word file whose embedded mammoth style map writes them, and a zip
of an HTML page and that Word file.

`numbered_report.docx` holds Word lists numbered other than 1, 2, 3 from the
top of each, which mammoth wrote as lists that did (docs/09 P31).

`definitions_book.epub` holds definition lists with terms in a row and a
definition after one of more than one block, which Python-Markdown read as
paragraphs and as text inside the definition before (docs/09 P64);
`definitions_archive.zip` holds an HTML page of the same and the book.

The two `quoted_lists_*` fixtures hold lists in a quotation in a list item,
with text or another list after them, which markdownify wrote as the item's
own: the text ran into the list's last item, and the second list was that
item's text or more of its items (docs/09 P74).  `quoted_lists_archive.zip`
holds the book's chapter as an HTML page.

The three `line_break_*` fixtures hold a line break (`<br>`, or Word's `w:br`)
in each place text flows: a paragraph, bold and italic text and a link, a list
item, a quotation, a table cell, and before text that reads as Markdown where
a line begins; and in the places it stays what it was: a heading, and code
(docs/09 P77).

The three `bare_link_*` fixtures hold links whose text is their target — a
file name, a path, a host, an HTML tag's name, an https address — in a
paragraph, a heading, a list item, a table cell and code (docs/09 P75).
"""

from __future__ import annotations

import base64
import csv
import html
import io
import json
import plistlib
import re
import struct
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURE_ROOT = REPO_ROOT / "tests" / "fixtures" / "office"

#: The strings the smoke test looks for. One per fixture, unmistakable in a
#: grep, and never written anywhere but into the document body.
SENTINELS = {
    "sample_report.docx": "SENTINEL DOCX WORD MARMOSET",
    "sample_deck.pptx": "SENTINEL PPTX SLIDE PANGOLIN",
    "sample_workbook.xlsx": "SENTINEL XLSX CELL AXOLOTL",
    "sample_book.epub": "SENTINEL EPUB CHAPTER NARWHAL",
    "numbered_book.epub": "SENTINEL EPUB NUMBERED KINKAJOU",
    "sample_archive.zip": "SENTINEL ZIP MEMBER OKAPI",
    # Inside the text box PowerPoint wraps in `mc:AlternateContent`, so it
    # reaches the bundle only if that box does.
    "equation_deck.pptx": "SENTINEL PPTX EQUATION QUOKKA",
    "equation_report.docx": "SENTINEL DOCX EQUATION WOMBAT",
    # The deck's own, from the deck inside it: in the wrapped text box, so it
    # reaches the bundle only if the rewrite reaches a zip's members too.
    "equation_archive.zip": "SENTINEL PPTX EQUATION QUOKKA",
    "markup_deck.pptx": "SENTINEL PPTX MARKUP CASSOWARY",
    "markup_report.docx": "SENTINEL DOCX MARKUP DUGONG",
    "markup_workbook.xlsx": "SENTINEL XLSX MARKUP GECKO",
    "markup_book.epub": "SENTINEL EPUB MARKUP IBEX",
    "markup_archive.zip": "SENTINEL ZIP MARKUP JERBOA",
    "left_out_archive.zip": "SENTINEL ZIP KEPT KINKAJOU",
    "adjacent_lists_book.epub": "SENTINEL EPUB LISTS TAMARIN",
    "picture_deck.pptx": "SENTINEL PPTX PICTURE KOALA",
    "picture_report.docx": "SENTINEL DOCX PICTURE LEMUR",
    "picture_book.epub": "SENTINEL EPUB PICTURE MANATEE",
    # Each ends its file, past MarkItDown's 64 KB sample, so it reaches the
    # bundle only if the file is decoded as the UTF-8 it is.
    "late_utf8.csv": "SENTINEL CSV LATE \N{LATIN CAPITAL LETTER N WITH TILDE}AND\N{LATIN CAPITAL LETTER U WITH ACUTE}",
    "late_utf8_archive.zip": "SENTINEL ZIP LATE P\N{LATIN CAPITAL LETTER E WITH ACUTE}CARI",
    "nested_book.epub": "SENTINEL EPUB NESTED CAPYBARA",
    "nested_report.docx": "SENTINEL DOCX NESTED BINTURONG",
    "nested_archive.zip": "SENTINEL ZIP NESTED FOSSA",
    "code_items_book.epub": "SENTINEL EPUB CODE GENET",
    "code_items_report.docx": "SENTINEL DOCX CODE HYRAX",
    "code_items_archive.zip": "SENTINEL ZIP CODE ICHNEUMON",
    "definition_book.epub": "SENTINEL EPUB DEFINITION MARGAY",
    "definition_archive.zip": "SENTINEL ZIP DEFINITION SERVAL",
    "inline_code_book.epub": "SENTINEL EPUB INLINE NUMBAT",
    "inline_code_report.docx": "SENTINEL DOCX INLINE OLINGO",
    "inline_code_archive.zip": "SENTINEL ZIP INLINE QUOLL",
    "quoted_code_book.epub": "SENTINEL EPUB QUOTED JACANA",
    "quoted_code_report.docx": "SENTINEL DOCX QUOTED LORIS",
    "quoted_code_archive.zip": "SENTINEL ZIP QUOTED MARGAY",
    "edge_space_report.docx": "SENTINEL DOCX EDGE SPACE LORIS",
    "edge_space_book.epub": "SENTINEL EPUB EDGE SPACE MARGAY",
    "edge_space_archive.zip": "SENTINEL ZIP EDGE SPACE NUMBAT",
    "whitespace_run_report.docx": "SENTINEL DOCX WHITESPACE RUN SERVAL",
    "whitespace_run_book.epub": "SENTINEL EPUB WHITESPACE RUN TAPIR",
    "whitespace_run_archive.zip": "SENTINEL ZIP WHITESPACE RUN VICUNA",
    "adjacent_blocks_book.epub": "SENTINEL EPUB BLOCKS LORIS",
    "adjacent_blocks_report.docx": "SENTINEL DOCX BLOCKS MARGAY",
    "adjacent_blocks_archive.zip": "SENTINEL ZIP BLOCKS NUMBAT",
    "numbered_report.docx": "SENTINEL DOCX NUMBERED TARSIER",
    "definitions_book.epub": "SENTINEL EPUB TERMS SERVAL",
    "definitions_archive.zip": "SENTINEL ZIP TERMS TARSIER",
    "quoted_lists_book.epub": "SENTINEL EPUB QUOTED ARMADILLO",
    "quoted_lists_archive.zip": "SENTINEL ZIP QUOTED TAPIR",
    "line_break_report.docx": "SENTINEL DOCX LINE BREAK WALLABY",
    "line_break_book.epub": "SENTINEL EPUB LINE BREAK XERUS",
    "line_break_archive.zip": "SENTINEL ZIP LINE BREAK ZORILLA",
    "bare_link_report.docx": "SENTINEL DOCX BARE LINK WALLABY",
    "bare_link_book.epub": "SENTINEL EPUB BARE LINK XERUS",
    "bare_link_archive.zip": "SENTINEL ZIP BARE LINK YAPOK",
}

#: Deterministic timestamp for every archive member, so regenerating a fixture
#: with no content change produces the same bytes (zip records mtimes).
FIXED_TIME = (2026, 1, 1, 0, 0, 0)


def _write_archive(
    path: Path, members: list[tuple[str, str | bytes]], *, first_stored: bool = False
) -> None:
    """Write a reproducible zip. `first_stored` is EPUB's uncompressed mimetype."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, (name, payload) in enumerate(members):
            info = zipfile.ZipInfo(name, date_time=FIXED_TIME)
            info.compress_type = (
                zipfile.ZIP_STORED
                if (first_stored and index == 0)
                else zipfile.ZIP_DEFLATED
            )
            info.external_attr = 0o644 << 16
            archive.writestr(info, payload)


# --------------------------------------------------------------------------- #
# DOCX — WordprocessingML
# --------------------------------------------------------------------------- #

_DOCX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>
"""

_DOCX_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""


def _docx_paragraph(text: str, style: str | None = None) -> str:
    properties = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{properties}<w:r><w:t xml:space='preserve'>{text}</w:t></w:r></w:p>"


def _docx_row(cells: list[str]) -> str:
    return (
        "<w:tr>"
        + "".join(f"<w:tc><w:tcPr/>{_docx_paragraph(cell)}</w:tc>" for cell in cells)
        + "</w:tr>"
    )


def write_docx(path: Path) -> None:
    body = (
        _docx_paragraph("Quarterly Field Report", "Heading1")
        + _docx_paragraph(SENTINELS["sample_report.docx"])
        + _docx_paragraph("Regional totals", "Heading2")
        + "<w:tbl><w:tblPr/>"
        + _docx_row(["Region", "Revenue"])
        + _docx_row(["North", "1250000"])
        + _docx_row(["South", "980000"])
        + "</w:tbl>"
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>\n"
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", _DOCX_CONTENT_TYPES),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
        ],
    )


# --------------------------------------------------------------------------- #
# EPUB — an OCF container around XHTML
# --------------------------------------------------------------------------- #

_EPUB_CONTAINER = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>
"""

_EPUB_OPF = """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="book-id">urn:uuid:webshot-office-fixture</dc:identifier>
    <dc:title>The Field Notebook</dc:title>
    <dc:creator>WebShot fixtures</dc:creator>
    <dc:language>en</dc:language>
  </metadata>
  <manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest>
  <spine><itemref idref="chapter"/></spine>
</package>
"""

_EPUB_CHAPTER = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter One</title></head>
<body>
  <h1>Chapter One</h1>
  <p>{SENTINELS["sample_book.epub"]}</p>
  <p>A second paragraph, so the chapter has more than a sentinel in it.</p>
</body></html>
"""


def write_epub(path: Path) -> None:
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", _EPUB_CHAPTER),
        ],
        first_stored=True,
    )


#: Lists that begin past 1, which markdownify writes as Markdown that begins
#: at that number (docs/09 P30). Each is closed by a paragraph: two adjacent
#: lists were one list to Python-Markdown until docs/09 P32.
_EPUB_NUMBERED_CHAPTER = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter Two</title></head>
<body>
  <h1>Chapter Two</h1>
  <p>{SENTINELS["numbered_book.epub"]}</p>
  <p>The survey resumes where the first chapter stopped.</p>
  <ol start="5"><li>Count the marsh plots</li><li>Weigh the samples</li></ol>
  <p>A list that begins at zero.</p>
  <ol start="0"><li>Calibrate the scale</li></ol>
  <p>A list inside a list.</p>
  <ol start="12"><li>Log the tide<ol start="4"><li>Note the wind</li></ol></li></ol>
  <p>A list that begins at one.</p>
  <ol><li>File the notebook</li></ol>
</body></html>
"""


def write_numbered_epub(path: Path) -> None:
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", _EPUB_NUMBERED_CHAPTER),
        ],
        first_stored=True,
    )


# --------------------------------------------------------------------------- #
# ZIP — a folder of ordinary documents
# --------------------------------------------------------------------------- #


def write_zip(path: Path) -> None:
    _write_archive(
        path,
        [
            (
                "notes/summary.txt",
                f"Field notes\n\n{SENTINELS['sample_archive.zip']}\n",
            ),
            ("data/regions.csv", "region,revenue\nNorth,1250000\nSouth,980000\n"),
        ],
    )


# --------------------------------------------------------------------------- #
# PPTX / XLSX — written with the readers' own libraries
# --------------------------------------------------------------------------- #


def write_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Field Review"
    slide.placeholders[1].text = SENTINELS["sample_deck.pptx"]
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))


def write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Regions"
    sheet.append(["Region", "Revenue", "Note"])
    sheet.append(["North", 1250000, SENTINELS["sample_workbook.xlsx"]])
    sheet.append(["South", 980000, "steady"])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))


# --------------------------------------------------------------------------- #
# Equations and what else the readers pass over (docs/09 P27)
# --------------------------------------------------------------------------- #

_OOXML = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "v": "urn:schemas-microsoft-com:vml",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
}
_DECLARE = " ".join(f'xmlns:{prefix}="{uri}"' for prefix, uri in _OOXML.items())

#: A 1x1 PNG: the picture PowerPoint fills the Fallback shape with.
_PIXEL_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360f80f0000010101005d0a2db4"
    "0000000049454e44ae426082"
)


def _math(inner: str) -> str:
    return f"<m:oMath>{inner}</m:oMath>"


def _math_run(text: str) -> str:
    # PowerPoint gives a math run DrawingML run properties; Word gives none.
    return f'<m:r><a:rPr lang="en-US" i="1"/><m:t>{text}</m:t></m:r>'


def _inline_equation(inner: str) -> str:
    return f"<a14:m>{_math(inner)}</a14:m>"


def _run(text: str, properties: str = "") -> str:
    return f"<a:r>{properties}<a:t>{text}</a:t></a:r>"


#: Fraction and delimiter properties are optional in the schema, and MarkItDown's
#: LaTeX converter needs them: these two are what it refuses.
_REFUSED_FRACTION = "<m:f><m:num><m:r><m:t>x</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>"
_REFUSED_DELIMITER = "<m:d><m:e><m:r><m:t>a+b</m:t></m:r></m:e></m:d>"


def _text_shape(
    shape_id: int, paragraphs: str, *, fill: str = "", y: int = 1600200
) -> str:
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="Text {shape_id}"/>'
        '<p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="457200" y="{y}"/><a:ext cx="8229600" cy="1500000"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill}</p:spPr>'
        f"<p:txBody><a:bodyPr/><a:lstStyle/>{paragraphs}</p:txBody></p:sp>"
    )


def _alternate(choice: str, fallback: str, requires: str = "a14") -> str:
    """What PowerPoint writes for a shape a 2007 reader cannot draw."""
    return (
        f'<mc:AlternateContent><mc:Choice Requires="{requires}">{choice}'
        f"</mc:Choice><mc:Fallback>{fallback}</mc:Fallback></mc:AlternateContent>"
    )


def _pictured(shape_id: int, picture: str, *, y: int = 1600200) -> str:
    """The Fallback: a no-break space over a picture of the real text."""
    fill = (
        f'<a:blipFill><a:blip r:embed="{picture}"/>'
        "<a:stretch><a:fillRect/></a:stretch></a:blipFill>"
    )
    return _text_shape(shape_id, f"<a:p>{_run(chr(0xA0))}</a:p>", fill=fill, y=y)


def _smartart_model() -> bytes:
    """Two steps under one diagram, stored in the opposite order to `srcOrd`."""
    point = (
        '<dgm:pt modelId="{{{id}}}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/>'
        "<a:lstStyle/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></dgm:t></dgm:pt>"
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<dgm:dataModel {_DECLARE}><dgm:ptLst>"
        '<dgm:pt modelId="{0}" type="doc"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/>'
        "<a:lstStyle/><a:p/></dgm:t></dgm:pt>"
        + point.format(id=2, text="Then weigh the second step")
        + point.format(id=1, text="First gather the readings")
        + '<dgm:pt modelId="{3}" type="pres"><dgm:prSet presName="diagram"/><dgm:spPr/></dgm:pt>'
        "</dgm:ptLst><dgm:cxnLst>"
        '<dgm:cxn modelId="{8}" srcId="{0}" destId="{2}" srcOrd="1" destOrd="0"/>'
        '<dgm:cxn modelId="{9}" srcId="{0}" destId="{1}" srcOrd="0" destOrd="0"/>'
        '<dgm:cxn modelId="{7}" type="presOf" srcId="{1}" destId="{3}" srcOrd="0" destOrd="0"/>'
        "</dgm:cxnLst><dgm:bg/><dgm:whole/></dgm:dataModel>"
    ).encode()


def _repack(data: bytes, path: Path) -> None:
    """python-pptx's archive again, with fixed timestamps so it regenerates byte for byte."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members: list[tuple[str, str | bytes]] = [
            (info.filename, archive.read(info)) for info in archive.infolist()
        ]
    _write_archive(path, members)


def write_equation_deck(path: Path) -> None:
    from pptx import Presentation
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml import parse_xml

    def place(parent: Any, markup: str) -> None:
        """Append `markup` to an element of python-pptx's tree, namespaces declared."""
        for element in list(parse_xml(f"<wrapper {_DECLARE}>{markup}</wrapper>")):
            parent.append(element)

    presentation = Presentation()
    title_only = presentation.slide_layouts[5]

    # 1. The text box PowerPoint writes when one holds an inline equation.
    slide = presentation.slides.add_slide(title_only)
    slide.shapes.title.text = "Bounds on Two Readings"
    _, picture = slide.part.get_or_add_image_part(io.BytesIO(_PIXEL_PNG))
    subscript = '<a:rPr baseline="-25000"/>'
    body = (
        f"<a:p>{_run(SENTINELS['equation_deck.pptx'])}</a:p><a:p>"
        + _run("x ")
        + _inline_equation(_math_run("≤"))
        + _run(" 12; y")
        + _run("1", subscript)
        + _run(" ")
        + _inline_equation(_math_run(">"))
        + _run(" 7")
        + "</a:p>"
    )
    place(slide.shapes._spTree, _alternate(_text_shape(3, body), _pictured(3, picture)))
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Say the bound aloud "
    place(notes.paragraphs[0]._p, _inline_equation(_math_run("≥")))

    # 2. The same, inside a group.
    slide = presentation.slides.add_slide(title_only)
    slide.shapes.title.text = "A Grouped Bound"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 3000000, 3000000, 500000).text = "Beside the group"
    _, picture = slide.part.get_or_add_image_part(io.BytesIO(_PIXEL_PNG))
    grouped = f"<a:p>{_run('z ')}{_inline_equation(_math_run('≥'))}{_run(' 3')}</a:p>"
    place(
        group._element,
        _alternate(_text_shape(30, grouped, y=0), _pictured(30, picture, y=0)),
    )

    # 3. An equation in a shape python-pptx does read.
    slide = presentation.slides.add_slide(title_only)
    slide.shapes.title.text = "A Bare Equation"
    bare = f"<a:p>{_run('a ')}{_inline_equation(_math_run('≠'))}{_run(' b')}</a:p>"
    place(slide.shapes._spTree, _text_shape(4, bare))

    # 4. SmartArt: its words are in a part of their own.
    slide = presentation.slides.add_slide(title_only)
    slide.shapes.title.text = "A Two-Step Process"
    model = Part(
        PackURI("/ppt/diagrams/data1.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
        slide.part.package,
        _smartart_model(),
    )
    data = slide.part.relate_to(
        model,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData",
    )
    place(
        slide.shapes._spTree,
        '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="5" name="Diagram 4"/>'
        "<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>"
        '<p:xfrm><a:off x="457200" y="1600200"/><a:ext cx="8229600" cy="3000000"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{_OOXML["dgm"]}"><dgm:relIds r:dm="{data}" '
        'r:lo="" r:qs="" r:cs=""/></a:graphicData></a:graphic></p:graphicFrame>',
    )

    # 5. What still cannot be read: an equation the converter refuses, and an
    # embedded object. Then ink, which stays wrapped: unwrapped, python-pptx
    # fails on it. Both wrapped as PowerPoint wraps them.
    slide = presentation.slides.add_slide(title_only)
    slide.shapes.title.text = "What Stays Unread"
    refused = f"<a:p>{_run('sum ')}{_inline_equation(_REFUSED_DELIMITER)}</a:p>"
    place(slide.shapes._spTree, _text_shape(6, refused))
    _, picture = slide.part.get_or_add_image_part(io.BytesIO(_PIXEL_PNG))
    embedded = (
        '<p:oleObj name="Equation" r:id="{rid}" progId="Equation.3">{inner}</p:oleObj>'
    )
    place(
        slide.shapes._spTree,
        '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="7" name="Object 6"/>'
        "<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>"
        '<p:xfrm><a:off x="457200" y="3600000"/><a:ext cx="3000000" cy="1000000"/></p:xfrm>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/presentationml/2006/ole">'
        + _alternate(
            embedded.format(rid=picture, inner="<p:embed/>"),
            embedded.format(
                rid=picture,
                inner=f'<p:embed/><p:pic><p:nvPicPr><p:cNvPr id="0" name=""/><p:cNvPicPr/>'
                f'<p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="{picture}"/>'
                "<a:stretch><a:fillRect/></a:stretch></p:blipFill><p:spPr/></p:pic>",
            ),
            requires="v",
        )
        + "</a:graphicData></a:graphic></p:graphicFrame>",
    )
    ink = (
        f'<p:contentPart p14:bwMode="auto" r:id="{picture}"><p14:nvContentPartPr>'
        '<p14:cNvPr id="8" name="Ink 7"/><p14:cNvContentPartPr/><p14:nvPr/>'
        '</p14:nvContentPartPr><p14:xfrm><a:off x="457200" y="5000000"/>'
        '<a:ext cx="1000000" cy="500000"/></p14:xfrm></p:contentPart>'
    )
    place(
        slide.shapes._spTree,
        _alternate(ink, _pictured(8, picture, y=5000000), requires="p14"),
    )

    buffer = io.BytesIO()
    presentation.save(buffer)
    _repack(buffer.getvalue(), path)


def _word_run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _word_math_run(text: str) -> str:
    return f"<m:r><m:t>{text}</m:t></m:r>"


def write_equation_report(path: Path) -> None:
    sup = (
        "<m:sSup><m:sSupPr><m:ctrlPr/></m:sSupPr><m:e>"
        f"{_word_math_run('c')}</m:e><m:sup>{_word_math_run('2')}</m:sup></m:sSup>"
    )
    boxed = f"<w:txbxContent><w:p>{_word_run('box text ')}{_math(_REFUSED_DELIMITER)}</w:p></w:txbxContent>"
    text_box = (
        '<w:p><w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>'
        '<wp:anchor><wp:docPr id="1" name="Text Box 1"/><a:graphic>'
        f'<a:graphicData uri="{_OOXML["wps"]}"><wps:wsp><wps:txbx>{boxed}</wps:txbx>'
        "<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>"
        f"</mc:Choice><mc:Fallback><w:pict><v:shape><v:textbox>{boxed}</v:textbox>"
        "</v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>"
    )
    body = (
        _docx_paragraph("Equation Notes", "Heading1")
        + _docx_paragraph(SENTINELS["equation_report.docx"])
        + f"<w:p>{_word_run('inline a ')}{_math(_word_math_run('≠'))}{_word_run(' b')}</w:p>"
        + f"<w:p><m:oMathPara>{_math(_word_math_run('E=m') + sup)}</m:oMathPara></w:p>"
        + f"<w:p>{_word_run('half of x is ')}{_math(_REFUSED_FRACTION)}</w:p>"
        + text_box
        # A chart, which mammoth drops: named in a warning, not read.
        + '<w:p><w:r><w:drawing><wp:inline><wp:extent cx="1" cy="1"/>'
        '<wp:docPr id="2" name="Chart 1"/><a:graphic>'
        f'<a:graphicData uri="{_OOXML["c"]}"><c:chart r:id="rId9"/></a:graphicData>'
        "</a:graphic></wp:inline></w:drawing></w:r></w:p>"
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", _DOCX_CONTENT_TYPES),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
        ],
    )


#: Where `equation_archive.zip` puts the two equation fixtures: in folders, so
#: a warning's member name is a path and not a bare file name.
EQUATION_MEMBERS = {
    "equation_deck.pptx": "slides/equation_deck.pptx",
    "equation_report.docx": "reports/equation_report.docx",
}


def write_equation_archive(path: Path) -> None:
    """The two equation fixtures, byte for byte, inside a zip (docs/09 P29).

    The same bytes, so each member's page can be held to the file's own; written
    fresh rather than copied from the fixtures beside it, so this does not
    depend on the order the fixtures are written in.
    """
    with tempfile.TemporaryDirectory() as scratch:
        members: list[tuple[str, str | bytes]] = []
        for name, member in EQUATION_MEMBERS.items():
            written = Path(scratch) / name
            WRITERS[name](written)
            members.append((member, written.read_bytes()))
    _write_archive(path, members)


# --------------------------------------------------------------------------- #
# Left out — the members a zip's reader cannot put on the page (docs/09 P33)
# --------------------------------------------------------------------------- #

#: Where the AppleDouble members say their files were downloaded from: the
#: address a capture must not put on the page.  `.invalid` never resolves.
DOWNLOADED_FROM = "https://downloads.example.invalid/week-1/notes.zip"


def _apple_double(attributes: dict[str, bytes]) -> bytes:
    """An AppleDouble file as the Finder writes one into `__MACOSX/`.

    RFC 1740's header and two entries, Finder info and an empty resource fork,
    with Apple's `ATTR` block for extended attributes after the Finder info,
    which is where a file's quarantine record and download address are kept.
    """
    names = [name.encode() + b"\0" for name in attributes]
    entries_size = sum((11 + len(name) + 3) // 4 * 4 for name in names)
    attr_start = 50 + 32 + 2  # header and two entry descriptors, Finder info, pad
    data_start = attr_start + 36 + entries_size
    data_length = sum(map(len, attributes.values()))
    total = data_start + data_length
    entries = b""
    data = b""
    for name, value in zip(names, attributes.values(), strict=True):
        entry = struct.pack(">IIHB", data_start + len(data), len(value), 0, len(name))
        entry += name
        entries += entry + b"\0" * (-len(entry) % 4)
        data += value
    attr_header = b"ATTR" + struct.pack(
        ">IIIIIIIHH", 0, total, data_start, data_length, 0, 0, 0, 0, len(names)
    )
    header = struct.pack(">II16sH", 0x00051607, 0x00020000, b"Mac OS X        ", 2)
    header += struct.pack(">III", 9, 50, total - 50)  # Finder info, with ATTR
    header += struct.pack(">III", 2, total, 0)  # resource fork, empty
    return header + b"TEXTttxt" + b"\0" * 24 + b"\0\0" + attr_header + entries + data


def _flagged_encrypted(data: bytes, name: str) -> bytes:
    """`data` with `name`'s encryption bit set in both of its headers.

    zipfile writes no encrypted member, and an archive tool would write
    different bytes on every run.  A reader decides by this bit before it reads
    the member, and so does WebShot, so the member's bytes can stay as written.
    """
    patched = bytearray(data)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entry = archive.getinfo(name)
        offset = archive.start_dir
    patched[entry.header_offset + 6] |= 0x1
    while patched[offset : offset + 4] == b"PK\x01\x02":
        lengths = struct.unpack_from("<HHH", patched, offset + 28)
        if patched[offset + 46 : offset + 46 + lengths[0]] == name.encode():
            patched[offset + 8] |= 0x1
            return bytes(patched)
        offset += 46 + sum(lengths)
    raise ValueError(f"{name} is not in the central directory")


def write_left_out_archive(path: Path) -> None:
    """One member of each kind a zip's reader leaves out, beside one it reads.

    A directory entry, a format nothing reads, a PDF (MarkItDown reads one, but
    with a package no WebShot extra installs), a Word file that is not one, an
    encrypted member, and two AppleDouble members, one of which records a
    download address.  `report.docx` is read with the office extra and left
    out, with WebShot's own advice, without it.
    """
    report = _docx_paragraph("Left-out archive report", "Heading1")
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{report}</w:body></w:document>\n"
    )
    with tempfile.TemporaryDirectory() as scratch:
        docx = Path(scratch) / "report.docx"
        _write_archive(
            docx,
            [
                ("[Content_Types].xml", _DOCX_CONTENT_TYPES),
                ("_rels/.rels", _DOCX_RELS),
                ("word/document.xml", document),
            ],
        )
        report_bytes = docx.read_bytes()
    where_from = plistlib.dumps([DOWNLOADED_FROM], fmt=plistlib.FMT_BINARY)
    members: list[tuple[str, str | bytes]] = [
        ("notes/", b""),
        ("notes/kept.txt", f"Kept notes\n\n{SENTINELS['left_out_archive.zip']}\n"),
        ("report.docx", report_bytes),
        ("figure.bin", bytes(range(256))),
        ("scan.pdf", _tiny_pdf("Left-out scan")),
        ("broken.docx", b"PK\x03\x04" + bytes(range(256))),
        ("secret.txt", "Encrypted notes\n"),
        ("__MACOSX/", b""),
        ("__MACOSX/notes/", b""),
        (
            "__MACOSX/notes/._kept.txt",
            _apple_double(
                {
                    "com.apple.metadata:kMDItemWhereFroms": where_from,
                    "com.apple.quarantine": b"0083;00000000;Safari;",
                }
            ),
        ),
        ("__MACOSX/._report.docx", _apple_double({})),
    ]
    _write_archive(path, members)
    path.write_bytes(_flagged_encrypted(path.read_bytes(), "secret.txt"))


def _tiny_pdf(text: str) -> bytes:
    """A one-page PDF of `text`, its cross-reference table offsets computed."""
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for number, content in enumerate(objects, 1):
        offsets.append(len(body))
        body += b"%d 0 obj\n%s\nendobj\n" % (number, content)
    table = b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    return (
        body
        + b"xref\n0 %d\n0000000000 65535 f \n%s" % (len(objects) + 1, table)
        + b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1)
        + b"startxref\n%d\n%%%%EOF\n" % len(body)
    )


# --------------------------------------------------------------------------- #
# Markup — text a document holds that HTML or Markdown would read as structure
# --------------------------------------------------------------------------- #

#: Where an image would be fetched from if a document's text reached the page
#: as markup.  `.invalid` never resolves (RFC 2606), so a capture that tried
#: says so as a failed request, and nothing here can reach a real host.
BEACON = "https://tracker.invalid"

#: An alt text and a link target that close the Markdown syntax holding them
#: and open an image of their own.
BREAKOUT_ALT = f"a chart]({BEACON}/alt.png) ![x"
BREAKOUT_HREF = f"https://example.invalid/?q=)![x]({BEACON}/href.png"
#: A code block whose own line closes markdownify's three-backtick fence.
BREAKOUT_PRE = f'one\n```\n<img src="{BEACON}/pre.png">\n```\ntwo'


def markup_lines(where: str) -> list[str]:
    """Lines of text that are markup to HTML or to Python-Markdown."""
    return [
        f'<img src="{BEACON}/{where}.png"> and *stars*',
        "# Not a heading",
        f"[a link](https://example.invalid/{where}) and ![an image]({BEACON}/{where}-md.png)",
        "3. Third, as typed",
        "Total",
        "===",
        "Term",
        ": not a definition",
        "~~~",
        "Set {x}",
        "C:\\Users\\_x_ and a\\*b and __init__",
        "&copy; as typed",
    ]


def _without_clock(data: bytes) -> bytes:
    """A python-pptx or openpyxl package, with its clock readings fixed.

    Both stamp the save time into `docProps/core.xml`, and python-pptx embeds a
    chart's data as an XlsxWriter workbook that does the same, so the zip's
    timestamps are not the only thing that would change on every run.
    """
    stamp = re.compile(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)")
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            payload = source.read(info)
            if info.filename.endswith("docProps/core.xml"):
                payload = stamp.sub(rb"\g<1>2026-01-01T00:00:00Z\g<2>", payload)
            elif info.filename.endswith(".xlsx"):
                payload = _without_clock(payload)
            entry = zipfile.ZipInfo(info.filename, date_time=FIXED_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            target.writestr(entry, payload)
    return output.getvalue()


def write_markup_deck(path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = f'<img src="{BEACON}/title.png"> and *title*'
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(3))
    lines = [SENTINELS["markup_deck.pptx"], *markup_lines("slide")]
    # A blank line, then a tab: Markdown's indented code block.
    lines += ["", "\tIndented, not code"]
    box.text_frame.text = lines[0]
    for line in lines[1:]:
        box.text_frame.add_paragraph().text = line
    table = slide.shapes.add_table(
        2, 2, Inches(0.5), Inches(4.6), Inches(9), Inches(0.8)
    ).table
    for column, (head, cell) in enumerate(
        [("Header", f'<img src="{BEACON}/cell.png">'), ("Other", "# cell")]
    ):
        table.cell(0, column).text = head
        table.cell(1, column).text = cell
    data = CategoryChartData()
    data.categories = [f'<img src="{BEACON}/category.png">']
    data.add_series("<b>series</b>", (1,))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5),
        Inches(5.5),
        Inches(4),
        Inches(1.5),
        data,
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "[chart](https://example.invalid/chart) _x_"
    slide.notes_slide.notes_text_frame.text = f'<img src="{BEACON}/notes.png"> *notes*'
    buffer = io.BytesIO()
    presentation.save(buffer)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_without_clock(buffer.getvalue()))


_MARKUP_DOCX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Default Extension="png" ContentType="image/png"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>
"""


def write_markup_report(path: Path) -> None:
    from xml.sax.saxutils import escape, quoteattr

    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"'
        f' Target={quoteattr(BREAKOUT_HREF)} TargetMode="External"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"'
        ' Target="media/image1.png"/></Relationships>\n'
    )
    picture = (
        '<w:p><w:r><w:drawing><wp:inline><wp:extent cx="9525" cy="9525"/>'
        f'<wp:docPr id="1" name="Picture 1" descr={quoteattr(BREAKOUT_ALT)}/>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:nvPicPr><pic:cNvPr id="0" name="pixel.png"/><pic:cNvPicPr/></pic:nvPicPr>'
        '<pic:blipFill><a:blip r:embed="rId3"/></pic:blipFill><pic:spPr/></pic:pic>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
    )
    body = (
        _docx_paragraph(SENTINELS["markup_report.docx"])
        + "".join(_docx_paragraph(escape(line)) for line in markup_lines("word"))
        + f'<w:p><w:hyperlink r:id="rId2">{_word_run("a link out")}</w:hyperlink></w:p>'
        + picture
        + "<w:tbl><w:tblPr/>"
        + _docx_row(["Header", "Other"])
        + _docx_row([escape(f'<img src="{BEACON}/cell.png">'), "# cell"])
        + "</w:tbl>"
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", _MARKUP_DOCX_CONTENT_TYPES),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
            ("word/media/image1.png", _PIXEL_PNG),
        ],
    )


#: A sheet name is a heading MarkItDown writes itself.  Excel forbids
#: `[]:*?/\` in one, which still leaves HTML and most of Markdown.
MARKUP_SHEET = "<i>Q1 #2 _x_ & {y}"


def write_markup_workbook(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MARKUP_SHEET
    rows = [["Where", "Text"], ["sentinel", SENTINELS["markup_workbook.xlsx"]]]
    rows += [
        [f"line {number}", line] for number, line in enumerate(markup_lines("cell"))
    ]
    rows.append(["formula text", "=SUM(A1)"])
    for row in rows:
        sheet.append(row)
        # openpyxl takes a leading `=` for a formula; these cells hold the text.
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    buffer = io.BytesIO()
    workbook.save(buffer)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_without_clock(buffer.getvalue()))


#: The book's own picture, inline: a file beside it would not be beside the
#: page WebShot renders, and its failed request would read as the text's.
_PIXEL_URI = "data:image/png;base64," + base64.b64encode(_PIXEL_PNG).decode()


def write_markup_book(path: Path) -> None:
    from xml.sax.saxutils import escape, quoteattr

    title = escape(f"<b>Markup</b> *book* & {{{SENTINELS['markup_book.epub']}}}")
    opf = f"""<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="book-id">urn:uuid:webshot-markup-fixture</dc:identifier>
    <dc:title>{title}</dc:title>
    <dc:creator># Creator</dc:creator>
    <dc:description>{escape(f'<img src="{BEACON}/description.png">')}</dc:description>
    <dc:language>en</dc:language>
  </metadata>
  <manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest>
  <spine><itemref idref="chapter"/></spine>
</package>
"""
    paragraphs = "".join(f"<p>{escape(line)}</p>" for line in markup_lines("chapter"))
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter</title></head>
<body>
  <h1>Chapter One</h1>
  <p>{SENTINELS["markup_book.epub"]}</p>
  {paragraphs}
  <p>A sentence that ends robots!<a href="#note">1</a></p>
  <p><a href={quoteattr(BREAKOUT_HREF)}>a link out</a></p>
  <p><img src="{_PIXEL_URI}" alt={quoteattr(BREAKOUT_ALT)}/></p>
  <pre>{escape(BREAKOUT_PRE)}</pre>
  <p id="note">1. The note.</p>
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", opf),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


def _csv(rows: list[list[str]]) -> str:
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerows(rows)
    return buffer.getvalue()


def write_markup_archive(path: Path) -> None:
    lines = [SENTINELS["markup_archive.zip"], *markup_lines("member")]
    cells = [
        f'<img src="{BEACON}/csv.png">',
        "=SUM(A1)",
        "~approx",
        "Set {x}",
        "C:\\Users\\_x_",
    ]
    _write_archive(
        path,
        [
            ("notes/markup.txt", "\n".join(lines) + "\n"),
            ("data/markup.csv", _csv([["Cell"], *([cell] for cell in cells)])),
            ("data/markup.tsv", "Cell\tOther\n" + f'<img src="{BEACON}/tsv.png">\t#\n'),
            # A member's name is text the archive's reader writes as a heading.
            ("notes/<img src=https:tracker.invalid>.txt", "named\n"),
        ],
    )


# --------------------------------------------------------------------------- #
# Lists side by side (docs/09 P32)
# --------------------------------------------------------------------------- #

#: Each pair of lists, by what stands between them in the chapter.  Python-
#: Markdown read every pair but the last as one list, until the bridge kept
#: them apart; the last is a bulleted list and a numbered one, which it never
#: joined.  The words are made up.
ADJACENT_LISTS = {
    "Two numbered lists": (
        '<ol start="5"><li>Brallow</li><li>Crindet</li></ol>'
        "<ol><li>Dulvane</li><li>Fenwisk</li></ol>"
    ),
    "Two bulleted lists": (
        "<ul><li>Glorvet</li><li>Hustrel</li></ul><ul><li>Imbrisk</li></ul>"
    ),
    "Nothing that shows between them": (
        "<ol><li>Jomvet</li></ol>\n  <div></div><p></p><!-- a note --><br/>"
        "<ol><li>Kravish</li></ol>"
    ),
    "Each in a division of its own": (
        "<div><ol><li>Lomtrel</li></ol></div><div><ol><li>Morvisk</li></ol></div>"
    ),
    "Inside one quotation": (
        "<blockquote><ul><li>Nuvrel</li></ul><ul><li>Opvane</li></ul></blockquote>"
    ),
    "A bulleted list, then a numbered one": (
        "<ul><li>Plindex</li></ul><ol><li>Quorvel</li></ol>"
    ),
}


#: The book's lists as its page holds them, in order: each tag and its items.
ADJACENT_LISTS_PAGE = [
    ("ol", ["Brallow", "Crindet"]),
    ("ol", ["Dulvane", "Fenwisk"]),
    ("ul", ["Glorvet", "Hustrel"]),
    ("ul", ["Imbrisk"]),
    ("ol", ["Jomvet"]),
    ("ol", ["Kravish"]),
    ("ol", ["Lomtrel"]),
    ("ol", ["Morvisk"]),
    ("ul", ["Nuvrel"]),
    ("ul", ["Opvane"]),
    ("ul", ["Plindex"]),
    ("ol", ["Quorvel"]),
    ("ol", ["Rendisk"]),
    ("ol", ["Stumbrel"]),
]


def write_adjacent_lists_book(path: Path) -> None:
    """Pairs of lists with nothing but markup between them, in two chapters.

    The first chapter ends with a list and the second begins with one:
    MarkItDown converts each chapter on its own and joins them with a blank
    line, so the two met only once the book was whole.
    """
    sections = "".join(
        f"<h2>{heading}</h2>{lists}" for heading, lists in ADJACENT_LISTS.items()
    )
    first = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Lists</title></head>
<body>
  <h1>Lists Side by Side</h1>
  <p>{SENTINELS["adjacent_lists_book.epub"]}</p>
  {sections}
  <h2>Across two chapters</h2>
  <ol><li>Rendisk</li></ol>
</body></html>
"""
    second = """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>More</title></head>
<body>
  <ol><li>Stumbrel</li></ol>
  <p>The second chapter goes on after its list.</p>
</body></html>
"""
    opf = """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="book-id">urn:uuid:webshot-adjacent-lists-fixture</dc:identifier>
    <dc:title>Lists Side by Side</dc:title>
    <dc:language>en</dc:language>
  </metadata>
  <manifest>
    <item id="one" href="one.xhtml" media-type="application/xhtml+xml"/>
    <item id="two" href="two.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="one"/><itemref idref="two"/></spine>
</package>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", opf),
            ("OEBPS/one.xhtml", first),
            ("OEBPS/two.xhtml", second),
        ],
        first_stored=True,
    )


# --------------------------------------------------------------------------- #
# Pictures — held, convertible, undrawable, linked, and named but missing
# --------------------------------------------------------------------------- #

#: A picture's alt text, which a placeholder and a warning both quote.
PICTURE_ALTS = {
    "png": "A blue rectangle",
    "tiff": "A green rectangle, stored as TIFF",
    "metafile": "A red rectangle, stored as a metafile",
    "linked": "A picture the deck only links to",
    "missing": "A picture the book names but does not hold",
}


def _picture(rgb: tuple[int, int, int], kind: str) -> bytes:
    """A small solid rectangle, in `kind` (`PNG` or `TIFF`), the same each run."""
    from PIL import Image

    output = io.BytesIO()
    Image.new("RGB", (120, 80), rgb).save(output, kind)
    return output.getvalue()


#: The smallest Windows Metafile Pillow and python-pptx take for one: a
#: placeable header, the metafile header and its end record, and nothing
#: drawn.  Nothing here draws a metafile, so the type is what the fixture is
#: for.  Real decks hold WMF clip art and Equation Editor previews (docs/09 P36).
_METAFILE = (
    struct.pack("<IHhhhhHIH", 0x9AC6CDD7, 0, 0, 0, 1200, 800, 1440, 0, 0)
    + struct.pack("<HHHIHIH", 1, 9, 0x300, 12, 0, 3, 0)
    + struct.pack("<IH", 3, 0)
)


def write_picture_deck(path: Path) -> None:
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    presentation = Presentation()

    def picture(slide: Any, data: bytes, alt: str, left: float) -> Any:
        shape = slide.shapes.add_picture(
            io.BytesIO(data), Inches(left), Inches(2.5), Inches(2.5), Inches(1.7)
        )
        shape._element.nvPicPr.cNvPr.set("descr", alt)
        return shape

    held = presentation.slides.add_slide(presentation.slide_layouts[5])
    held.shapes.title.text = "Pictures the deck holds"
    box = held.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(0.6))
    box.text_frame.text = SENTINELS["picture_deck.pptx"]
    picture(held, _picture((20, 60, 200), "PNG"), PICTURE_ALTS["png"], 0.5)
    picture(held, _picture((20, 160, 60), "TIFF"), PICTURE_ALTS["tiff"], 3.5)

    unread = presentation.slides.add_slide(presentation.slide_layouts[5])
    unread.shapes.title.text = "Pictures the page cannot draw"
    picture(unread, _METAFILE, PICTURE_ALTS["metafile"], 0.5)
    # A linked picture: the blip names an external target and holds nothing.
    linked = picture(unread, _picture((90, 90, 90), "PNG"), PICTURE_ALTS["linked"], 3.5)
    blip = linked._element.blipFill.blip
    target = unread.part.relate_to(
        "https://example.invalid/linked.png", RELATIONSHIP_TYPE.IMAGE, is_external=True
    )
    del blip.attrib[qn("r:embed")]
    blip.set(qn("r:link"), target)

    buffer = io.BytesIO()
    presentation.save(buffer)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_without_clock(buffer.getvalue()))


def write_picture_report(path: Path) -> None:
    def drawing(number: int, alt: str) -> str:
        return (
            f'<w:p><w:r><w:drawing><wp:inline><wp:extent cx="1143000" cy="762000"/>'
            f'<wp:docPr id="{number}" name="Picture {number}" descr="{alt}"/>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<pic:nvPicPr><pic:cNvPr id="0" name="picture{number}"/><pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="rId{number}"/></pic:blipFill><pic:spPr/></pic:pic>'
            "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
        )

    media = [
        ("image1.png", _picture((20, 60, 200), "PNG"), PICTURE_ALTS["png"]),
        ("image2.tiff", _picture((20, 160, 60), "TIFF"), PICTURE_ALTS["tiff"]),
        ("image3.emf", _METAFILE, PICTURE_ALTS["metafile"]),
    ]
    body = _docx_paragraph(SENTINELS["picture_report.docx"]) + "".join(
        drawing(number, alt) for number, (_, _, alt) in enumerate(media, start=1)
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(
            f'<Relationship Id="rId{number}" Type="http://schemas.openxmlformats.org/'
            f'officeDocument/2006/relationships/image" Target="media/{name}"/>'
            for number, (name, _, _) in enumerate(media, start=1)
        )
        + "</Relationships>\n"
    )
    content_types = _DOCX_CONTENT_TYPES.replace(
        "  <Override",
        '  <Default Extension="png" ContentType="image/png"/>\n'
        '  <Default Extension="tiff" ContentType="image/tiff"/>\n'
        '  <Default Extension="emf" ContentType="image/x-emf"/>\n'
        "  <Override",
        1,
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", content_types),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
            *((f"word/media/{name}", data) for name, data, _ in media),
        ],
    )


def write_picture_book(path: Path) -> None:
    opf = """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="book-id">urn:uuid:webshot-picture-fixture</dc:identifier>
    <dc:title>The Picture Book</dc:title>
    <dc:language>en</dc:language>
  </metadata>
  <manifest>
    <item id="chapter" href="text/chapter.xhtml" media-type="application/xhtml+xml"/>
    <item id="blue" href="images/blue.png" media-type="image/png"/>
    <item id="green" href="images/green.tif" media-type="image/tiff"/>
  </manifest>
  <spine><itemref idref="chapter"/></spine>
</package>
"""
    # The chapter sits a folder down, so its pictures are named `../images/…`.
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter</title></head>
<body>
  <h1>Chapter One</h1>
  <p>{SENTINELS["picture_book.epub"]}</p>
  <p><img src="../images/blue.png" alt="{PICTURE_ALTS["png"]}"/></p>
  <p><img src="../images/green.tif" alt="{PICTURE_ALTS["tiff"]}"/></p>
  <p><img src="../images/missing.png" alt="{PICTURE_ALTS["missing"]}"/></p>
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", opf),
            ("OEBPS/text/chapter.xhtml", chapter),
            ("OEBPS/images/blue.png", _picture((20, 60, 200), "PNG")),
            ("OEBPS/images/green.tif", _picture((20, 160, 60), "TIFF")),
        ],
        first_stored=True,
    )


# --------------------------------------------------------------------------- #
# Text whose first character outside ASCII is past MarkItDown's sample
# --------------------------------------------------------------------------- #

#: Bytes of ASCII before each late character: past the 65,536 MarkItDown's
#: `_read_charset_sample` reads, with room to spare.
LATE_AFTER = 80_000
#: The late character's words in the members that do not carry the sentinel.
LATE_WORDS = (
    "\N{LATIN CAPITAL LETTER N WITH TILDE}and\N{LATIN SMALL LETTER U WITH ACUTE}"
)


def _repeated(line: str) -> str:
    """`line` as many times as it takes to fill `LATE_AFTER` bytes."""
    return line * (LATE_AFTER // len(line) + 1)


#: A long cell, so a table passes `LATE_AFTER` in few rows: capturing a
#: table costs time with the square of its rows (docling-core's DocTags
#: serializer rebuilds the table's grid once per row), and 11,431 short rows
#: took the round-trip test past 40 minutes.
LATE_CELL = " ".join(
    ["Counted on the morning walk, one line per stretch of the path"] * 16
)


def _late_rows(delimiter: str, last: str) -> str:
    """A header, numbered rows of `LATE_CELL` past `LATE_AFTER`, and `last`.

    Numbered, because Magika must take the file for CSV, and so for text,
    for MarkItDown to guess its encoding at all: ending in the sentinel after
    rows of `x,1`, it was a file Magika could not name.
    """
    count = LATE_AFTER // len(f"00{delimiter}{LATE_CELL}\n") + 1
    rows = "".join(f"{index}{delimiter}{LATE_CELL}\n" for index in range(count))
    return f"id{delimiter}notes\n{rows}{count}{delimiter}{last}\n"


def write_late_utf8_csv(path: Path) -> None:
    """docs/09 P35's measured case: on its own, it ended the run with exit 2."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_late_rows(",", SENTINELS["late_utf8.csv"]).encode("utf-8"))


def _late_notebook() -> str:
    # Prose, not code: charset_normalizer guessed UTF-8 for a sample of
    # `total = total + 1`, ASCII though it was, and the notebook read.
    cells = [
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": _repeated("Counting, one row at a time.\n").splitlines(
                keepends=True
            ),
        },
        {"cell_type": "markdown", "metadata": {}, "source": [f"# {LATE_WORDS}\n"]},
    ]
    notebook = {"cells": cells, "metadata": {}, "nbformat": 4, "nbformat_minor": 5}
    return json.dumps(notebook, ensure_ascii=False, indent=1) + "\n"


def write_late_utf8_archive(path: Path) -> None:
    """A member of each kind the late character cost, and three it did not.

    Read by its 64 KB guess, MarkItDown failed the Python, JSON, YAML,
    notebook and TSV members with `UnicodeDecodeError`.  Of the two short
    files behind a byte order mark, it found no reader for the YAML one, whose
    content Magika could not name, and put the mark on the page as U+FEFF for
    the text one, which Magika named.  `windows.csv` is Windows-1252, which
    WebShot read only as UTF-8 until docs/09 P38, and now reads as the
    windows-1252 it guesses, and says so.  Every word is made up.
    """
    module = (
        '"""Generated values."""\n\n\n'
        + "".join(
            f"def value_{index}(x):\n    return x + {index}\n\n\n"
            for index in range(LATE_AFTER // 40 + 1)
        )
        + f'NAME = "{SENTINELS["late_utf8_archive.zip"]}"\n'
    )
    members: list[tuple[str, str | bytes]] = [
        ("src/values.py", module),
        (
            "data/records.json",
            "[\n"
            + _repeated('  {"id": 1, "name": "x"},\n')
            + f'  {{"id": 2, "name": "{LATE_WORDS}"}}\n]\n',
        ),
        (
            "data/records.yaml",
            "items:\n"
            + _repeated("  - name: x\n    id: 1\n")
            + f"  - name: {LATE_WORDS}\n",
        ),
        ("notebooks/count.ipynb", _late_notebook()),
        ("data/rows.tsv", _late_rows("\t", LATE_WORDS)),
        ("data/marked.yaml", "\ufeffname: " + LATE_WORDS + "\n"),
        (
            "notes/marked.txt",
            "\ufeffNotes\n\nThe first line of a note, written in a text editor "
            "that marks its encoding.\nA second line follows it, and a third.\n",
        ),
        ("data/windows.csv", f"a,b\n{LATE_WORDS},1\n".encode("cp1252")),
    ]
    _write_archive(path, members)


# --------------------------------------------------------------------------- #
# Lists nested in lists (docs/09 P40)
# --------------------------------------------------------------------------- #

#: A list as the page should hold it: its tag, and each item's own text with
#: the lists nested in it.  The words are made up.
NestedList = tuple[str, list[tuple[str, list["NestedList"]]]]

#: The book's lists, by what each section shows.  markdownify indented every
#: nested list by its parent's marker width, and Python-Markdown nests only at
#: four spaces, so each of these but the quotation lost its nesting.
NESTED_BOOK_SECTIONS = {
    "Bullets under bullets": (
        "<ul><li>Arvel<ul><li>Bostrin</li><li>Calvane</li></ul></li>"
        "<li>Dempor</li></ul>"
    ),
    "Three levels, each kind in turn": (
        "<ul><li>Elvorn<ol><li>Fanrow<ul><li>Gesh</li></ul></li>"
        "<li>Hollet</li></ol></li></ul>"
    ),
    "Numbers under numbers": (
        "<ol><li>Imbrel<ol><li>Jastow<ol><li>Kelvorn</li></ol></li></ol></li>"
        "<li>Lumbret</li></ol>"
    ),
    "Numbers of one digit and of two": (
        '<ol start="9"><li>Mivash<ul><li>Norvel</li></ul></li>'
        "<li>Orbisk<ul><li>Pellow</li></ul></li></ol>"
    ),
    "Items of more than one paragraph": (
        "<ol><li><p>Quorvane</p><p>Rasket</p></li><li><p>Sulvane</p></li></ol>"
    ),
    "A bulleted list and a numbered one in one item": (
        "<ol><li>Tornish<ul><li>Umbrel</li></ul><ol><li>Vasket</li></ol></li>"
        "<li>Wendrel</li></ol>"
    ),
    "Two numbered lists in one item": (
        "<ul><li>Zelvane<ol><li>Abrisk</li></ol><ol><li>Bentrow</li></ol></li></ul>"
    ),
    "Text after a nested list": (
        "<ul><li>Corvel<ul><li>Dastrin</li></ul>and Elbrisk after it</li>"
        "<li>Fennick</li></ul>"
    ),
    "An item that is only a list": (
        "<ul><li><ul><li>Gorvane</li><li>Holvish</li></ul></li><li>Istrel</li></ul>"
    ),
    "A quotation in an item": (
        "<ul><li>Jorvel<blockquote><p>Kestrin</p></blockquote></li></ul>"
    ),
}

#: `NESTED_BOOK_SECTIONS`' lists as the page holds them, in order.
NESTED_BOOK_PAGE: list[NestedList] = [
    ("ul", [("Arvel", [("ul", [("Bostrin", []), ("Calvane", [])])]), ("Dempor", [])]),
    (
        "ul",
        [("Elvorn", [("ol", [("Fanrow", [("ul", [("Gesh", [])])]), ("Hollet", [])])])],
    ),
    (
        "ol",
        [
            ("Imbrel", [("ol", [("Jastow", [("ol", [("Kelvorn", [])])])])]),
            ("Lumbret", []),
        ],
    ),
    (
        "ol",
        [
            ("Mivash", [("ul", [("Norvel", [])])]),
            ("Orbisk", [("ul", [("Pellow", [])])]),
        ],
    ),
    ("ol", [("Quorvane Rasket", []), ("Sulvane", [])]),
    (
        "ol",
        [
            ("Tornish", [("ul", [("Umbrel", [])]), ("ol", [("Vasket", [])])]),
            ("Wendrel", []),
        ],
    ),
    ("ul", [("Zelvane", [("ol", [("Abrisk", [])]), ("ol", [("Bentrow", [])])])]),
    (
        "ul",
        [("Corvel and Elbrisk after it", [("ul", [("Dastrin", [])])]), ("Fennick", [])],
    ),
    ("ul", [("", [("ul", [("Gorvane", []), ("Holvish", [])])]), ("Istrel", [])]),
    ("ul", [("Jorvel Kestrin", [])]),
]


def write_nested_book(path: Path) -> None:
    sections = "".join(
        f"<h2>{heading}</h2>{lists}" for heading, lists in NESTED_BOOK_SECTIONS.items()
    )
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Nested</title></head>
<body>
  <h1>Lists Within Lists</h1>
  <p>{SENTINELS["nested_book.epub"]}</p>
  {sections}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


#: Word's numbering: an outline numbered 1., a., i. by level, and bullets.
_NESTED_FORMATS = {
    1: ("decimal", "lowerLetter", "lowerRoman"),
    2: ("bullet", "bullet", "bullet"),
}

#: The report's list paragraphs: (numbering, level, text), and None for a
#: paragraph of body text, which ends a list.  The second list is the shape
#: mammoth writes as a bulleted list and a numbered one in a single item.
NESTED_REPORT_PARAGRAPHS: list[tuple[int, int, str] | None] = [
    (1, 0, "Lorvane"),
    (1, 1, "Mestrin"),
    (1, 1, "Nolvish"),
    (1, 2, "Ospret"),
    (1, 0, "Pindrel"),
    None,
    (1, 0, "Quellan"),
    (2, 1, "Rostrel"),
    (1, 1, "Sabrine"),
    (1, 0, "Tolvane"),
    None,
    (2, 0, "Umbrisk"),
    (2, 1, "Vellow"),
    (2, 2, "Wistrel"),
    (2, 0, "Yorvane"),
]

NESTED_REPORT_PAGE: list[NestedList] = [
    (
        "ol",
        [
            (
                "Lorvane",
                [("ol", [("Mestrin", []), ("Nolvish", [("ol", [("Ospret", [])])])])],
            ),
            ("Pindrel", []),
        ],
    ),
    (
        "ol",
        [
            ("Quellan", [("ul", [("Rostrel", [])]), ("ol", [("Sabrine", [])])]),
            ("Tolvane", []),
        ],
    ),
    (
        "ul",
        [
            ("Umbrisk", [("ul", [("Vellow", [("ul", [("Wistrel", [])])])])]),
            ("Yorvane", []),
        ],
    ),
]

_NESTED_DOCX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>
</Types>
"""


def write_nested_report(path: Path) -> None:
    abstracts = "".join(
        f'<w:abstractNum w:abstractNumId="{number}">'
        + "".join(
            f'<w:lvl w:ilvl="{level}"><w:start w:val="1"/>'
            f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="%{level + 1}."/></w:lvl>'
            for level, fmt in enumerate(formats)
        )
        + "</w:abstractNum>"
        for number, formats in _NESTED_FORMATS.items()
    )
    instances = "".join(
        f'<w:num w:numId="{number}"><w:abstractNumId w:val="{number}"/></w:num>'
        for number in _NESTED_FORMATS
    )
    numbering = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<w:numbering xmlns:w="{_OOXML["w"]}">{abstracts}{instances}</w:numbering>\n'
    )
    items = "".join(
        _docx_paragraph("The lists go on below.")
        if paragraph is None
        else (
            f'<w:p><w:pPr><w:numPr><w:ilvl w:val="{paragraph[1]}"/>'
            f'<w:numId w:val="{paragraph[0]}"/></w:numPr></w:pPr>'
            f"{_word_run(paragraph[2])}</w:p>"
        )
        for paragraph in NESTED_REPORT_PARAGRAPHS
    )
    body = (
        _docx_paragraph("Lists Within Lists", "Heading1")
        + _docx_paragraph(SENTINELS["nested_report.docx"])
        + items
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering"'
        ' Target="numbering.xml"/></Relationships>\n'
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", _NESTED_DOCX_CONTENT_TYPES),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
            ("word/numbering.xml", numbering),
        ],
    )


#: An HTML page in the archive: a zip member is read by the same writer.
_NESTED_MEMBER_PAGE = f"""<!DOCTYPE html>
<html><head><title>Steps</title></head>
<body>
  <p>{SENTINELS["nested_archive.zip"]}</p>
  <ul><li>Zandrel<ol><li>Arbisk<ul><li>Belvane</li></ul></li></ol></li></ul>
</body></html>
"""

#: What the archive's page holds: its own page's list, then the report's.
NESTED_ARCHIVE_PAGE: list[NestedList] = [
    ("ul", [("Zandrel", [("ol", [("Arbisk", [("ul", [("Belvane", [])])])])])]),
    *NESTED_REPORT_PAGE,
]


def write_nested_archive(path: Path) -> None:
    """An HTML page with a nested list, and the nested report as it is."""
    with tempfile.TemporaryDirectory() as scratch:
        report = Path(scratch) / "nested_report.docx"
        write_nested_report(report)
        _write_archive(
            path,
            [
                ("notes/steps.html", _NESTED_MEMBER_PAGE),
                ("reports/nested_report.docx", report.read_bytes()),
            ],
        )


#: The book's items, by what each section shows: a definition list in a list
#: item, whose definitions hold more than a paragraph.  Python-Markdown read
#: every block after a definition's first into the item, so each section but
#: the first-paragraph-only parts lost its definition's structure.  The words
#: are made up.
DEFINITION_BOOK_SECTIONS = {
    "A definition's second paragraph": (
        "<ul><li>Arbelt<dl><dt>Brosken</dt><dd><p>Cavrel</p><p>Dunmorrow</p>"
        "</dd></dl></li><li>Elvast</li></ul>"
    ),
    "A numbered list in a definition": (
        "<ol><li>Farnish<dl><dt>Gorvel</dt><dd>Hastrel<ol><li>Ilvane</li>"
        "<li>Jorvish</li></ol></dd></dl></li><li>Kembrisk</li></ol>"
    ),
    "A bulleted list two deep in a definition": (
        "<ul><li>Lumvar<dl><dt>Mostrel</dt><dd>Nevish<ul><li>Orbane<ul>"
        "<li>Pelvish</li></ul></li></ul></dd></dl></li></ul>"
    ),
    "A definition list in a definition": (
        "<ol><li>Quarvel<dl><dt>Rostrin</dt><dd>Selvane<dl><dt>Tormish</dt>"
        "<dd><p>Ulbrin</p><p>Vastrow</p></dd></dl></dd></dl></li></ol>"
    ),
    "A quotation in a definition": (
        "<ul><li>Wendrisk<dl><dt>Yarbel</dt><dd>Zolvane<blockquote>"
        "<p>Arnwick</p><p>Belstrow</p></blockquote></dd></dl></li></ul>"
    ),
    "Three items deep": (
        "<ol><li>Corbish<ul><li>Delvar<ol><li>Emberwick<dl><dt>Falstrin</dt>"
        "<dd><p>Gemvar</p><p>Holbrisk</p></dd></dl></li></ol></li></ul></li></ol>"
    ),
    "An item that opens with a definition list": (
        "<ul><li><dl><dt>Ivrenn</dt><dd><ol><li>Jalbrin</li><li>Korvash</li>"
        "</ol></dd></dl></li><li>Lendrow</li></ul>"
    ),
    "The item's own text after its definition list": (
        "<ol><li>Marvisk<dl><dt>Norbel</dt><dd><p>Olvish</p><p>Parnow</p></dd>"
        "</dl><p>Quilvane</p></li></ol>"
    ),
}


def write_definition_book(path: Path) -> None:
    sections = "".join(
        f"<h2>{heading}</h2>{body}"
        for heading, body in DEFINITION_BOOK_SECTIONS.items()
    )
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Terms</title></head>
<body>
  <h1>Terms Within Lists</h1>
  <p>{SENTINELS["definition_book.epub"]}</p>
  {sections}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


#: The archive's HTML page's own item: a definition of two paragraphs and a
#: list, in a numbered item.
DEFINITION_MEMBER_ITEMS = (
    "<ol><li>Rendlow<dl><dt>Sarvish</dt><dd><p>Talbrin</p><p>Umbrisk</p>"
    "<ul><li>Velnor</li></ul></dd></dl></li><li>Wistrel</li></ol>"
)

_DEFINITION_MEMBER_PAGE = f"""<!DOCTYPE html>
<html><head><title>Glossary</title></head>
<body>
  <p>{SENTINELS["definition_archive.zip"]}</p>
  {DEFINITION_MEMBER_ITEMS}
</body></html>
"""


def write_definition_archive(path: Path) -> None:
    """An HTML page with a definition list in an item, and the book as it is."""
    with tempfile.TemporaryDirectory() as scratch:
        book = Path(scratch) / "definition_book.epub"
        write_definition_book(book)
        _write_archive(
            path,
            [
                ("notes/glossary.html", _DEFINITION_MEMBER_PAGE),
                ("books/definition_book.epub", book.read_bytes()),
            ],
        )


# --------------------------------------------------------------------------- #
# Code in list items (docs/09 P44)
# --------------------------------------------------------------------------- #

#: A code block as the page should hold it: how many items it sits in (0 at
#: the top level), the text of the item that holds it but for its code and its
#: nested lists, and the code.  The words are made up.
ItemCode = tuple[int, str, str]

#: The book's code, by what each section shows.  markdownify fenced each
#: `<pre>` and the item indented the fence, which Python-Markdown reads only
#: where a line begins: every block but the last became one line of text.
CODE_ITEMS_BOOK_SECTIONS = {
    "Code under a bullet": (
        "<ul><li>Arbel<pre>plinth = gorvane(7)\n  wex plinth</pre></li>"
        "<li>Bostick</li></ul>"
    ),
    "Code two levels down": (
        "<ol><li>Cestor<ol><li>Dravik<pre>hollow(3)\n    tamber</pre></li>"
        "<li>Elkin</li></ol></li><li>Forrel</li></ol>"
    ),
    "Code three levels down": (
        "<ul><li>Gantry<ol><li>Hesk<ul><li>Ivrel<pre>quill\n  quorn\n    quiver"
        "</pre></li></ul></li></ol></li><li>Jabrel</li></ul>"
    ),
    "An item that opens with code": (
        "<ul><li><pre>opal = 1\nopine = 2</pre></li><li>Jessop</li></ul>"
    ),
    "Code after a nested list": (
        "<ol><li>Kelvane<ul><li>Lurrow</li></ul><pre>after lurrow\n  still kelvane"
        "</pre></li><li>Mordel</li></ol>"
    ),
    "Two blocks of code in one item": (
        "<ul><li>Nostrel<pre>first block</pre><pre>second block</pre></li>"
        "<li>Ostler</li></ul>"
    ),
    "Blank lines, indentation, tabs, backticks and markup": (
        "<ul><li>Orvash<pre>\n   lead\n\nmid\n\tTabbed\n```\nfence\n````\n"
        "``x`` &lt;b&gt;not bold&lt;/b&gt; &amp;amp;\n</pre>and Pemberly after it"
        "</li><li>Quell</li></ul>"
    ),
    "Numbers of one digit and of two": (
        '<ol start="9"><li>Rendel<pre>nine</pre></li><li>Sorvel<pre>ten</pre></li></ol>'
    ),
    "Code outside a list": "<pre>top level\n  unchanged</pre>",
}

#: Every `<pre>` in `CODE_ITEMS_BOOK_SECTIONS`, as the page holds it, in order.
#: The tab is four spaces: Python-Markdown expands tabs, at the top level too.
CODE_ITEMS_BOOK_CODE: list[ItemCode] = [
    (1, "Arbel", "plinth = gorvane(7)\n  wex plinth"),
    (2, "Dravik", "hollow(3)\n    tamber"),
    (3, "Ivrel", "quill\n  quorn\n    quiver"),
    (1, "", "opal = 1\nopine = 2"),
    (1, "Kelvane", "after lurrow\n  still kelvane"),
    (1, "Nostrel", "first block"),
    (1, "Nostrel", "second block"),
    (
        1,
        "Orvash and Pemberly after it",
        "   lead\n\nmid\n    Tabbed\n```\nfence\n````\n``x`` <b>not bold</b> &amp;",
    ),
    (1, "Rendel", "nine"),
    (1, "Sorvel", "ten"),
    (0, "", "top level\n  unchanged"),
]

#: `CODE_ITEMS_BOOK_SECTIONS`' lists as the page holds them, code left out.
CODE_ITEMS_BOOK_PAGE: list[NestedList] = [
    ("ul", [("Arbel", []), ("Bostick", [])]),
    ("ol", [("Cestor", [("ol", [("Dravik", []), ("Elkin", [])])]), ("Forrel", [])]),
    (
        "ul",
        [
            ("Gantry", [("ol", [("Hesk", [("ul", [("Ivrel", [])])])])]),
            ("Jabrel", []),
        ],
    ),
    ("ul", [("", []), ("Jessop", [])]),
    ("ol", [("Kelvane", [("ul", [("Lurrow", [])])]), ("Mordel", [])]),
    ("ul", [("Nostrel", []), ("Ostler", [])]),
    ("ul", [("Orvash and Pemberly after it", []), ("Quell", [])]),
    ("ol", [("Rendel", []), ("Sorvel", [])]),
]


def write_code_items_book(path: Path) -> None:
    sections = "".join(
        f"<h2>{heading}</h2>{body}"
        for heading, body in CODE_ITEMS_BOOK_SECTIONS.items()
    )
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Code</title></head>
<body>
  <h1>Code Within Lists</h1>
  <p>{SENTINELS["code_items_book.epub"]}</p>
  {sections}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


#: mammoth writes a `<pre>` only where a style map says to.  A Word file can
#: carry one, and MarkItDown reads it: code paragraphs, each style a level of
#: list item deep, and consecutive ones one block.
CODE_ITEMS_STYLE_MAP = (
    "p.ItemCode => ul|ol > li > pre:separator('\\n')\n"
    "p.SubItemCode => ul|ol > li > ul|ol > li > pre:separator('\\n')\n"
)

#: The report's paragraphs: (level, text) for an item of a list numbered at
#: level 0 and bulleted at level 1; (style, text) for a line of code; None for
#: a paragraph of body text, which ends a list.
CODE_ITEMS_REPORT_PARAGRAPHS: list[tuple[int | str, str] | None] = [
    (0, "Vellum"),
    ("ItemCode", "wendle(3)"),
    ("ItemCode", "    yarrow = 2"),
    (1, "Wistan"),
    ("SubItemCode", "zephyr()"),
    ("SubItemCode", "  ablet"),
    (0, "Xandrel"),
    None,
    (0, "Yolvane"),
    (1, "Zamber"),
    ("ItemCode", "after zamber"),
    (0, "Abernot"),
]

CODE_ITEMS_REPORT_CODE: list[ItemCode] = [
    (1, "Vellum", "wendle(3)\n    yarrow = 2"),
    (2, "Wistan", "zephyr()\n  ablet"),
    (1, "Yolvane", "after zamber"),
]

CODE_ITEMS_REPORT_PAGE: list[NestedList] = [
    ("ol", [("Vellum", [("ul", [("Wistan", [])])]), ("Xandrel", [])]),
    ("ol", [("Yolvane", [("ul", [("Zamber", [])])]), ("Abernot", [])]),
]

_CODE_ITEMS_NUMBERING = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    f'<w:numbering xmlns:w="{_OOXML["w"]}">'
    '<w:abstractNum w:abstractNumId="1">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    '<w:lvlText w:val="%1."/></w:lvl>'
    '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
    '<w:lvlText w:val="-"/></w:lvl>'
    "</w:abstractNum>"
    '<w:num w:numId="1"><w:abstractNumId w:val="1"/></w:num>'
    "</w:numbering>\n"
)


def write_code_items_report(path: Path) -> None:
    """A Word file whose embedded style map puts code in its list items.

    Laid out as mammoth's own `embed_style_map` writes one: the map at
    `mammoth/style-map`, named in the content types and the document's
    relationships.
    """
    paragraphs = []
    for paragraph in CODE_ITEMS_REPORT_PARAGRAPHS:
        if paragraph is None:
            paragraphs.append(_docx_paragraph("The lists go on below."))
        elif isinstance(paragraph[0], str):
            paragraphs.append(_docx_paragraph(paragraph[1], paragraph[0]))
        else:
            paragraphs.append(
                f'<w:p><w:pPr><w:numPr><w:ilvl w:val="{paragraph[0]}"/>'
                f'<w:numId w:val="1"/></w:numPr></w:pPr>'
                f"{_word_run(paragraph[1])}</w:p>"
            )
    body = (
        _docx_paragraph("Code Within Lists", "Heading1")
        + _docx_paragraph(SENTINELS["code_items_report.docx"])
        + "".join(paragraphs)
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    content_types = _NESTED_DOCX_CONTENT_TYPES.replace(
        "</Types>",
        '  <Override PartName="/mammoth/style-map"'
        ' ContentType="text/prs.mammoth.style-map"/>\n</Types>',
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering"'
        ' Target="numbering.xml"/>'
        '<Relationship Id="rMammothStyleMap" Type="http://schemas.zwobble.org/mammoth/style-map"'
        ' Target="/mammoth/style-map"/></Relationships>\n'
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", content_types),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
            ("word/numbering.xml", _CODE_ITEMS_NUMBERING),
            ("mammoth/style-map", CODE_ITEMS_STYLE_MAP),
        ],
    )


#: An HTML page in the archive: a zip member is read by the same writer.
_CODE_ITEMS_MEMBER_PAGE = f"""<!DOCTYPE html>
<html><head><title>Snippets</title></head>
<body>
  <p>{SENTINELS["code_items_archive.zip"]}</p>
  <ol><li>Bramble<ul><li>Cindrel<pre>cindrel --go\n  --far</pre></li></ul></li>
  <li>Dunmore</li></ol>
</body></html>
"""

#: What the archive's page holds: its own page's, then the report's.
CODE_ITEMS_ARCHIVE_CODE: list[ItemCode] = [
    (2, "Cindrel", "cindrel --go\n  --far"),
    *CODE_ITEMS_REPORT_CODE,
]
CODE_ITEMS_ARCHIVE_PAGE: list[NestedList] = [
    ("ol", [("Bramble", [("ul", [("Cindrel", [])])]), ("Dunmore", [])]),
    *CODE_ITEMS_REPORT_PAGE,
]


def write_code_items_archive(path: Path) -> None:
    """An HTML page with code in a nested item, and the code report as it is."""
    with tempfile.TemporaryDirectory() as scratch:
        report = Path(scratch) / "code_items_report.docx"
        write_code_items_report(report)
        _write_archive(
            path,
            [
                ("notes/snippets.html", _CODE_ITEMS_MEMBER_PAGE),
                ("reports/code_items_report.docx", report.read_bytes()),
            ],
        )


# --------------------------------------------------------------------------- #
# Code in a line of text (docs/09 P59)
# --------------------------------------------------------------------------- #

#: Code as the page should hold it in a line of text: the element its
#: `<code>` sits in, and its text, each `<br>` a line break.
InlineCode = tuple[str, str]


def _beacon(spot: str) -> str:
    """An `<img>` typed in code, aimed at a `.invalid` host named for its spot."""
    return f'<img src="{BEACON}/inline-{spot}.png">'


#: Code that markdownify wrote into a line of text, each holding what took
#: the page there: a blank line, or a line Python-Markdown reads wherever a
#: paragraph's line begins — a heading, a rule, a quotation, a `<div>`.
INLINE_LINK = f"gorvel(2)\n\n{_beacon('link')}"
INLINE_STRONG = f"drest = 4\n\n![tamsin]({BEACON}/inline-strong.png)"
INLINE_EM = f"pell\n# not a heading\n{_beacon('em')}"
INLINE_DEL = f"quarrow\n***\n{_beacon('del')}"
INLINE_SUB = f"sennet\n> not quoted\n{_beacon('sub')}"
INLINE_Q = f"tolly\n<div>{_beacon('q')}</div>"
INLINE_CODE = f"ulvar\n\n{_beacon('code')}"
INLINE_LINES = f"harl = 2\n# not a heading\n{_beacon('lines')}"

#: The book's code, by what each section shows.
INLINE_CODE_BOOK_SECTIONS = {
    "Code in a link": (
        '<p>Read <a href="chapter.xhtml"><pre>'
        f"{html.escape(INLINE_LINK)}</pre></a> first.</p>"
    ),
    "Code in strong emphasis": (
        f"<p><strong><pre>{html.escape(INLINE_STRONG)}</pre></strong></p>"
    ),
    "Code in emphasis, in an item": (
        f"<ul><li>Quessel <em><pre>{html.escape(INLINE_EM)}</pre></em></li>"
        "<li>Rellick</li></ul>"
    ),
    "Code in struck text": f"<p><del><pre>{html.escape(INLINE_DEL)}</pre></del></p>",
    "Code in a subscript, after a word": (
        f"<p>Wrenfold<sub><pre>{html.escape(INLINE_SUB)}</pre></sub>tail</p>"
    ),
    "Code in a quotation": (
        f"<p>Ostrel said <q><pre>{html.escape(INLINE_Q)}</pre></q></p>"
    ),
    "Code in inline code": (
        f"<p><code><pre>{html.escape(INLINE_CODE)}</pre></code></p>"
    ),
    "Inline code over three lines": (
        f"<p><code>{html.escape(INLINE_LINES)}</code></p>"
    ),
    "Code with no blank line, in a link": (
        '<p><a href="chapter.xhtml"><pre>plain one\n  plain two</pre></a></p>'
    ),
    "Code outside a line of text": "<pre>top level\n  unchanged</pre>",
}

#: Every `<code>` the page holds outside a code block, in order.  Python-
#: Markdown has no struck text, so `~~` is the paragraph's; a quotation and a
#: subscript are written as their text.  In inline code a line break is a
#: space, as the browser shows one there.
INLINE_CODE_BOOK_CODE: list[InlineCode] = [
    ("a", INLINE_LINK),
    ("strong", INLINE_STRONG),
    ("em", INLINE_EM),
    ("p", INLINE_DEL),
    ("p", INLINE_SUB),
    ("p", INLINE_Q),
    ("p", INLINE_CODE.replace("\n", " ")),
    ("p", INLINE_LINES.replace("\n", " ")),
    ("a", "plain one\n  plain two"),
]
#: Every code block the page holds, which a `<pre>` outside a line keeps.
INLINE_CODE_BOOK_BLOCKS = ["top level\n  unchanged"]


def write_inline_code_book(path: Path) -> None:
    sections = "".join(
        f"<h2>{heading}</h2>{body}"
        for heading, body in INLINE_CODE_BOOK_SECTIONS.items()
    )
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Inline</title></head>
<body>
  <h1>Code Within Text</h1>
  <p>{SENTINELS["inline_code_book.epub"]}</p>
  {sections}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


#: A style map that writes code paragraphs into strong text after the body
#: text before them, and into emphasis in a list item.  mammoth drops an
#: empty paragraph, so these reach the page by a line, not a blank one.
INLINE_CODE_STYLE_MAP = (
    "p.StrongCode => p > strong > pre:separator('\\n')\n"
    "p.ItemEmCode => ul > li > em > pre:separator('\\n')\n"
)
INLINE_REPORT_STRONG = f"wendle(3)\n***\n{_beacon('report-strong')}"
INLINE_REPORT_EM = f"zephyr()\n# not a heading\n{_beacon('report-em')}"
INLINE_CODE_REPORT_CODE: list[InlineCode] = [
    ("strong", INLINE_REPORT_STRONG),
    ("em", INLINE_REPORT_EM),
]


def write_inline_code_report(path: Path) -> None:
    """A Word file whose embedded style map puts code into a line of text."""
    paragraphs = [
        _docx_paragraph("Code Within Text", "Heading1"),
        _docx_paragraph(SENTINELS["inline_code_report.docx"]),
        _docx_paragraph("Vantor writes"),
        *(
            _docx_paragraph(html.escape(line, quote=False), "StrongCode")
            for line in INLINE_REPORT_STRONG.split("\n")
        ),
        _docx_paragraph("And in a list:"),
        *(
            _docx_paragraph(html.escape(line, quote=False), "ItemEmCode")
            for line in INLINE_REPORT_EM.split("\n")
        ),
    ]
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{''.join(paragraphs)}</w:body></w:document>\n"
    )
    content_types = _DOCX_CONTENT_TYPES.replace(
        "</Types>",
        '  <Override PartName="/mammoth/style-map"'
        ' ContentType="text/prs.mammoth.style-map"/>\n</Types>',
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rMammothStyleMap" Type="http://schemas.zwobble.org/mammoth/style-map"'
        ' Target="/mammoth/style-map"/></Relationships>\n'
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", content_types),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
            ("mammoth/style-map", INLINE_CODE_STYLE_MAP),
        ],
    )


INLINE_MEMBER_LINK = f"brisk --run\n\n{_beacon('member-link')}"
INLINE_MEMBER_ITEM = f"cobble\n***\n{_beacon('member-item')}"
INLINE_MEMBER_KBD = f"dorrit\n> not quoted\n{_beacon('member-kbd')}"

#: An HTML page in the archive: a zip member is read by the same writer.
_INLINE_CODE_MEMBER_PAGE = f"""<!DOCTYPE html>
<html><head><title>Snippets</title></head>
<body>
  <p>{SENTINELS["inline_code_archive.zip"]}</p>
  <p><a href="#top"><pre>{html.escape(INLINE_MEMBER_LINK)}</pre></a></p>
  <ol><li>Bramwell <strong><pre>{html.escape(INLINE_MEMBER_ITEM)}</pre></strong></li>
  <li>Corvid</li></ol>
  <p>Press <kbd>{html.escape(INLINE_MEMBER_KBD)}</kbd> to go on.</p>
</body></html>
"""

#: What the archive's page holds: its own page's, then the report's.
INLINE_CODE_ARCHIVE_CODE: list[InlineCode] = [
    ("a", INLINE_MEMBER_LINK),
    ("strong", INLINE_MEMBER_ITEM),
    ("p", INLINE_MEMBER_KBD.replace("\n", " ")),
    *INLINE_CODE_REPORT_CODE,
]


def write_inline_code_archive(path: Path) -> None:
    """An HTML page with code in a line of text, and the inline code report."""
    with tempfile.TemporaryDirectory() as scratch:
        report = Path(scratch) / "inline_code_report.docx"
        write_inline_code_report(report)
        _write_archive(
            path,
            [
                ("notes/snippets.html", _INLINE_CODE_MEMBER_PAGE),
                ("reports/inline_code_report.docx", report.read_bytes()),
            ],
        )


# --------------------------------------------------------------------------- #
# Code in quotations and definitions (docs/09 P47)
# --------------------------------------------------------------------------- #

#: A code block as the page should hold it: the quotations (`blockquote`),
#: definitions (`dd`) and list items (`li`) it sits in, outermost first and
#: joined by `/` ("" at the top level), and the code.  The words are made up.
QuotedCode = tuple[str, str]

#: Code holding blank lines, indentation, a tab, backtick runs and markup:
#: markdownify fenced it, and behind `> ` or `:   ` the fence was text, so
#: the lines after a blank line were read as Markdown with raw HTML.
_QUOTED_MARKUP_CODE = (
    "\n   lead\n\nmid\n\tTabbed\n```\nfence\n````\n``x`` &lt;b&gt;not bold&lt;/b&gt;"
    ' &amp;amp;\n\n&lt;img src="https://tracker.invalid/quoted.png"&gt;\n'
)
#: `_QUOTED_MARKUP_CODE` as the page holds it: its blank lines at either end
#: trimmed, its tab four spaces, as Python-Markdown expands tabs everywhere.
_QUOTED_MARKUP_PAGE = (
    "   lead\n\nmid\n    Tabbed\n```\nfence\n````\n``x`` <b>not bold</b> &amp;\n\n"
    '<img src="https://tracker.invalid/quoted.png">'
)

#: The book's code, by what each section shows.  Before, every block but the
#: last was one line of text, or two paragraphs where it held a blank line.
QUOTED_CODE_BOOK_SECTIONS = {
    "Code in a quotation": (
        "<blockquote><p>Quillon wrote</p><pre>hollin(2)\n  vesper</pre>"
        "<p>Rastel replied</p></blockquote>"
    ),
    "A quotation of code alone": "<blockquote><pre>tamsk = 4\n    ulmer</pre></blockquote>",
    "Code in a quotation in a quotation": (
        "<blockquote><p>Sorrel</p><blockquote><p>Tibbet</p>"
        "<pre>nested(1)\n  deeper</pre></blockquote><p>Umber</p></blockquote>"
    ),
    "Code three quotations deep": (
        "<blockquote><blockquote><blockquote><pre>third\n  level</pre>"
        "</blockquote></blockquote></blockquote>"
    ),
    "Two blocks of code in a quotation": (
        "<blockquote><pre>first wend</pre><pre>second wend</pre></blockquote>"
    ),
    "Code after a list in a quotation": (
        "<blockquote><ul><li>Vantor</li></ul><pre>after vantor\n  still quoted</pre>"
        "</blockquote>"
    ),
    "Code in a definition": (
        "<dl><dt>Wexley</dt><dd><pre>wex --all\n  --quiet</pre></dd>"
        "<dt>Yarrow</dt><dd>Plain zell</dd></dl>"
    ),
    "Code after a definition's text": (
        "<dl><dt>Abbet</dt><dd>Brindle<pre>brindle()\n  cog</pre>"
        "and Corvy after it</dd></dl>"
    ),
    "Code in a definition in a quotation": (
        "<blockquote><dl><dt>Dunlin</dt><dd><pre>dunlin = 5\n  fen</pre></dd></dl>"
        "</blockquote>"
    ),
    "Code in a quotation in a definition": (
        "<dl><dt>Ember</dt><dd><blockquote><pre>ember\n  glow</pre></blockquote>"
        "</dd></dl>"
    ),
    "Blank lines, indentation, tabs, backticks and markup in a quotation": (
        f"<blockquote><pre>{_QUOTED_MARKUP_CODE}</pre></blockquote>"
    ),
    "Blank lines, indentation, tabs, backticks and markup in a definition": (
        f"<dl><dt>Farrow</dt><dd><pre>{_QUOTED_MARKUP_CODE}</pre></dd></dl>"
    ),
    "Code outside a quotation": "<pre>top level\n  unchanged</pre>",
}

#: Every `<pre>` in `QUOTED_CODE_BOOK_SECTIONS`, as the page holds it, in order.
QUOTED_CODE_BOOK_CODE: list[QuotedCode] = [
    ("blockquote", "hollin(2)\n  vesper"),
    ("blockquote", "tamsk = 4\n    ulmer"),
    ("blockquote/blockquote", "nested(1)\n  deeper"),
    ("blockquote/blockquote/blockquote", "third\n  level"),
    ("blockquote", "first wend"),
    ("blockquote", "second wend"),
    ("blockquote", "after vantor\n  still quoted"),
    ("dd", "wex --all\n  --quiet"),
    ("dd", "brindle()\n  cog"),
    ("blockquote/dd", "dunlin = 5\n  fen"),
    ("dd/blockquote", "ember\n  glow"),
    ("blockquote", _QUOTED_MARKUP_PAGE),
    ("dd", _QUOTED_MARKUP_PAGE),
    ("", "top level\n  unchanged"),
]


def write_quoted_code_book(path: Path) -> None:
    sections = "".join(
        f"<h2>{heading}</h2>{body}"
        for heading, body in QUOTED_CODE_BOOK_SECTIONS.items()
    )
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Quoted</title></head>
<body>
  <h1>Code Within Quotations</h1>
  <p>{SENTINELS["quoted_code_book.epub"]}</p>
  {sections}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


#: A style for a quotation's paragraphs and its code, a definition list's
#: terms, text and code, and code in a definition in a quotation: consecutive
#: code paragraphs are one block.
QUOTED_CODE_STYLE_MAP = (
    "p.Quote => blockquote > p:fresh\n"
    "p.QuoteCode => blockquote > pre:separator('\\n')\n"
    "p.Term => dl > dt:fresh\n"
    "p.Definition => dl > dd:fresh > p:fresh\n"
    "p.DefinitionCode => dl > dd > pre:separator('\\n')\n"
    "p.QuoteTerm => blockquote > dl > dt:fresh\n"
    "p.QuoteDefinitionCode => blockquote > dl > dd > pre:separator('\\n')\n"
)

#: The report's paragraphs, each its style and its text; a style of None is
#: body text, which ends a quotation or a definition list.
QUOTED_CODE_REPORT_PARAGRAPHS: list[tuple[str | None, str]] = [
    ("Quote", "Fennick wrote"),
    ("QuoteCode", "gorse(3)"),
    ("QuoteCode", "    hazel = 2"),
    ("Quote", "and Fennick went on"),
    (None, "The definitions follow."),
    ("Term", "Ilsey"),
    ("DefinitionCode", "jorum()"),
    ("DefinitionCode", "  kestle"),
    ("Term", "Lathom"),
    ("Definition", "Mallow text"),
    ("DefinitionCode", "after mallow"),
    (None, "A quoted definition follows."),
    ("QuoteTerm", "Nimbel"),
    ("QuoteDefinitionCode", "nimbel --on"),
    ("QuoteDefinitionCode", "  --oriel"),
]

QUOTED_CODE_REPORT_CODE: list[QuotedCode] = [
    ("blockquote", "gorse(3)\n    hazel = 2"),
    ("dd", "jorum()\n  kestle"),
    ("dd", "after mallow"),
    ("blockquote/dd", "nimbel --on\n  --oriel"),
]


def write_quoted_code_report(path: Path) -> None:
    """A Word file whose embedded style map puts code in quotations and
    definitions, laid out as `write_code_items_report`'s is."""
    body = (
        _docx_paragraph("Code Within Quotations", "Heading1")
        + _docx_paragraph(SENTINELS["quoted_code_report.docx"])
        + "".join(
            _docx_paragraph(text, style)
            for style, text in QUOTED_CODE_REPORT_PARAGRAPHS
        )
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    content_types = _DOCX_CONTENT_TYPES.replace(
        "</Types>",
        '  <Override PartName="/mammoth/style-map"'
        ' ContentType="text/prs.mammoth.style-map"/>\n</Types>',
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rMammothStyleMap" Type="http://schemas.zwobble.org/mammoth/style-map"'
        ' Target="/mammoth/style-map"/></Relationships>\n'
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", content_types),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
            ("mammoth/style-map", QUOTED_CODE_STYLE_MAP),
        ],
    )


#: An HTML page in the archive: a zip member is read by the same writer.
_QUOTED_CODE_MEMBER_PAGE = f"""<!DOCTYPE html>
<html><head><title>Quoted snippets</title></head>
<body>
  <p>{SENTINELS["quoted_code_archive.zip"]}</p>
  <blockquote><p>Oxley</p><blockquote><pre>oxley(1)\n  pell</pre></blockquote>
  </blockquote>
  <dl><dt>Quarrel</dt><dd><p>Rook</p><pre>quarrel = 7\n  sedge</pre></dd></dl>
</body></html>
"""

#: What the archive's page holds: its own page's, then the report's.
QUOTED_CODE_ARCHIVE_CODE: list[QuotedCode] = [
    ("blockquote/blockquote", "oxley(1)\n  pell"),
    ("dd", "quarrel = 7\n  sedge"),
    *QUOTED_CODE_REPORT_CODE,
]


def write_quoted_code_archive(path: Path) -> None:
    """An HTML page with code in a quotation and a definition, and the quoted
    code report as it is."""
    with tempfile.TemporaryDirectory() as scratch:
        report = Path(scratch) / "quoted_code_report.docx"
        write_quoted_code_report(report)
        _write_archive(
            path,
            [
                ("notes/quoted.html", _QUOTED_CODE_MEMBER_PAGE),
                ("reports/quoted_code_report.docx", report.read_bytes()),
            ],
        )


# --------------------------------------------------------------------------- #
# Edge spaces — a Unicode space at the edge of formatted text (docs/09 P43)
# --------------------------------------------------------------------------- #

#: Spaces outside ASCII that `str.strip` takes from the edge of a run, one of
#: each kind a document is likely to hold.  `tests/test_edge_spaces.py` derives
#: the full set from Python itself; these are the ones the fixtures carry.
EDGE_SPACES = {
    "nbsp": "\N{NO-BREAK SPACE}",
    "ensp": "\N{EN SPACE}",
    "thinsp": "\N{THIN SPACE}",
    "nnbsp": "\N{NARROW NO-BREAK SPACE}",
    "ideosp": "\N{IDEOGRAPHIC SPACE}",
}


@dataclass(frozen=True)
class EdgeSpaceLine:
    """One paragraph of an `edge_space_*` fixture: three made-up words, the
    middle one formatted, and every space between them exactly as written.

    `opening` and `closing` are the markup around the middle word: HTML tags
    for a book or a page, the names in `_WORD_RUN_PROPERTIES` for Word.
    """

    opening: str
    closing: str
    before: str
    lead: str
    inside: str
    trail: str
    after: str

    @property
    def run(self) -> str:
        """The formatted run's own text, its edge spaces included."""
        return self.lead + self.inside + self.trail

    @property
    def text(self) -> str:
        """The paragraph as the page must read it, character for character."""
        return self.before + self.run + self.after


def _edge_space_line(
    opening: str, closing: str, label: str, space: str, where: str
) -> EdgeSpaceLine:
    """A paragraph whose words are kept apart only by `space`, at `where`.

    Every word is `label` and `where` with a suffix, so each is unique in its
    fixture and a failure names the case.  An edge without `space` has an
    ASCII space outside the run instead; `only` is a run of the space alone.
    """
    stem = f"{label}{where}"
    if where == "only":
        return EdgeSpaceLine(
            opening, closing, f"{stem}before", space, "", "", f"{stem}after"
        )
    lead = space if where in ("lead", "both") else ""
    trail = space if where in ("trail", "both") else ""
    return EdgeSpaceLine(
        opening,
        closing,
        f"{stem}before" + ("" if lead else " "),
        lead,
        f"{stem}inside",
        trail,
        ("" if trail else " ") + f"{stem}after",
    )


def _edge_space_lines(kinds: dict[str, tuple[str, str]]) -> list[EdgeSpaceLine]:
    """Each kind with a no-break space at each edge, each other space at both
    edges of the first kind, and two lines whose output must not move: ASCII
    spaces at the run's edges, and spaces outside it."""
    lines = [
        _edge_space_line(opening, closing, label, EDGE_SPACES["nbsp"], where)
        for label, (opening, closing) in kinds.items()
        for where in ("lead", "trail", "both", "only")
    ]
    first, (opening, closing) = next(iter(kinds.items()))
    lines += [
        _edge_space_line(opening, closing, f"{first}{name}", space, "both")
        for name, space in EDGE_SPACES.items()
        if name != "nbsp"
    ]
    lines.append(_edge_space_line(opening, closing, f"{first}ascii", " ", "both"))
    lines.append(_edge_space_line(opening, closing, f"{first}outside", "", "both"))
    return lines


#: The formatting Word gives a run, by the element mammoth writes for it.
_WORD_RUN_PROPERTIES = {
    "bold": "<w:b/>",
    "italic": "<w:i/>",
    "bolditalic": "<w:b/><w:i/>",
    "strike": "<w:strike/>",
    "subscript": '<w:vertAlign w:val="subscript"/>',
    "superscript": '<w:vertAlign w:val="superscript"/>',
    "link": "",
}
#: Where the Word fixture's links point; `.invalid` never resolves.
EDGE_SPACE_HREF = "https://example.invalid/edge"
#: Every inline element markdownify `chomp`s, and two it reaches through them:
#: `strike` is `s`, and `a` without an `href` is text.
_HTML_INLINE = {
    "strong": ("<strong>", "</strong>"),
    "b": ("<b>", "</b>"),
    "em": ("<em>", "</em>"),
    "i": ("<i>", "</i>"),
    "nested": ("<strong><em>", "</em></strong>"),
    "code": ("<code>", "</code>"),
    "kbd": ("<kbd>", "</kbd>"),
    "samp": ("<samp>", "</samp>"),
    "del": ("<del>", "</del>"),
    "s": ("<s>", "</s>"),
    "strike": ("<strike>", "</strike>"),
    "sub": ("<sub>", "</sub>"),
    "sup": ("<sup>", "</sup>"),
    "u": ("<u>", "</u>"),
    "link": (f'<a href="{EDGE_SPACE_HREF}">', "</a>"),
    "anchor": ('<a id="edge">', "</a>"),
}


def edge_space_lines(name: str) -> list[EdgeSpaceLine]:
    """The paragraphs of the `edge_space_*` fixture called `name`."""
    if name == "edge_space_report.docx":
        return _edge_space_lines({kind: (kind, "") for kind in _WORD_RUN_PROPERTIES})
    return _edge_space_lines(_HTML_INLINE)


def _word_edge_paragraph(line: EdgeSpaceLine) -> str:
    def plain(text: str) -> str:
        return _word_run(text) if text else ""

    properties = _WORD_RUN_PROPERTIES[line.opening]
    run = f'<w:r><w:rPr>{properties}</w:rPr><w:t xml:space="preserve">{line.run}</w:t></w:r>'
    if line.opening == "link":
        run = f'<w:hyperlink r:id="rId2">{_word_run(line.run)}</w:hyperlink>'
    return f"<w:p>{plain(line.before)}{run}{plain(line.after)}</w:p>"


def write_edge_space_report(path: Path) -> None:
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"'
        f' Target="{EDGE_SPACE_HREF}" TargetMode="External"/></Relationships>\n'
    )
    body = _docx_paragraph(SENTINELS["edge_space_report.docx"]) + "".join(
        _word_edge_paragraph(line)
        for line in edge_space_lines("edge_space_report.docx")
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", _DOCX_CONTENT_TYPES),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
        ],
    )


def _edge_space_html(name: str) -> str:
    return "\n".join(
        f"  <p>{line.before}{line.opening}{line.run}{line.closing}{line.after}</p>"
        for line in edge_space_lines(name)
    )


def write_edge_space_book(path: Path) -> None:
    opf = _EPUB_OPF.replace("webshot-office-fixture", "webshot-edge-space-fixture")
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter One</title></head>
<body>
  <p>{SENTINELS["edge_space_book.epub"]}</p>
{_edge_space_html("edge_space_book.epub")}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", opf),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


def write_edge_space_archive(path: Path) -> None:
    """An HTML page in a zip: the one way a page reaches MarkItDown's HTML
    reader, since WebShot opens an `.html` file in the browser itself."""
    page = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Edge spaces</title></head>
<body>
  <p>{SENTINELS["edge_space_archive.zip"]}</p>
{_edge_space_html("edge_space_archive.zip")}
</body></html>
"""
    _write_archive(path, [("pages/edge-spaces.html", page)])


# --------------------------------------------------------------------------- #
# Whitespace runs — ASCII whitespace or a line break in formatted text (P49)
# --------------------------------------------------------------------------- #

#: What a formatted run holds, as HTML, and whether a plain space stands
#: before and after it.  `{inside}` is a made-up word; a space stands outside
#: the run only at an edge the run does not fill, so that only what the run
#: holds keeps the words apart.  The `space` edges were written right before
#: P49, except in an `<a>` MarkItDown writes without its markup.
_WHITESPACE_CASES = {
    "space": (" ", False, False),
    "spaces": ("  ", False, False),
    "tab": ("\t", False, False),
    "newline": ("\n", False, False),
    "br": ("<br/>", False, False),
    "leadnewline": ("\n{inside}", False, True),
    "trailnewline": ("{inside}\n", True, False),
    "leadbr": ("<br/>{inside}", False, True),
    "trailbr": ("{inside}<br/>", True, False),
    "leadspace": (" {inside}", False, True),
    "trailspace": ("{inside} ", True, False),
}
#: The same runs as Word holds them, inside one `w:r`: Word has no newline in
#: a run's text, and writes a tab and a line break as elements of their own.
_WORD_WHITESPACE_CASES = {
    "space": ('<w:t xml:space="preserve"> </w:t>', False, False),
    "spaces": ('<w:t xml:space="preserve">  </w:t>', False, False),
    "tab": ("<w:tab/>", False, False),
    "br": ("<w:br/>", False, False),
    "leadbr": ("<w:br/><w:t>{inside}</w:t>", False, True),
    "trailbr": ("<w:t>{inside}</w:t><w:br/>", True, False),
    "leadspace": ('<w:t xml:space="preserve"> {inside}</w:t>', False, True),
    "trailspace": ('<w:t xml:space="preserve">{inside} </w:t>', True, False),
}
#: A link whose text is its address, which MarkItDown wrote as `<address>`
#: until the bridge turned that form off (docs/09 P75).
WHITESPACE_RUN_HREF = "https://example.invalid/whitespace"


@dataclass(frozen=True)
class WhitespaceRunLine:
    """One paragraph of a `whitespace_run_*` fixture.

    `run` is the formatted run as the document holds it: HTML for a book or a
    page, a `w:r` or `w:hyperlink` for Word.  `before` and `after` are the
    plain text either side.  `words` are what the page must read, with
    whitespace between each and nowhere else.
    """

    before: str
    run: str
    after: str
    words: tuple[str, ...]

    @property
    def text(self) -> str:
        """The paragraph as a reader reads it: its words, one space apart."""
        return " ".join(self.words)


def _whitespace_line(
    label: str, content: str, space_before: bool, space_after: bool, wrap: Any
) -> WhitespaceRunLine:
    """A paragraph `{label}prior`, a formatted run, `{label}next`.

    `content` is what the run holds; its word, where it has one, is
    `{label}core`.  `wrap` puts the content inside the run's formatting.
    """
    prior, core, following = f"{label}prior", f"{label}core", f"{label}next"
    words = (prior, core, following) if "{inside}" in content else (prior, following)
    return WhitespaceRunLine(
        prior + (" " if space_before else ""),
        wrap(content.replace("{inside}", core)),
        (" " if space_after else "") + following,
        words,
    )


def whitespace_run_lines(name: str) -> list[WhitespaceRunLine]:
    """The paragraphs of the `whitespace_run_*` fixture called `name`."""
    if name == "whitespace_run_report.docx":
        return [
            _whitespace_line(
                f"{kind}{case}",
                content,
                space_before,
                space_after,
                lambda inner, kind=kind: _word_whitespace_run(kind, inner),
            )
            for kind in _WORD_RUN_PROPERTIES
            for case, (
                content,
                space_before,
                space_after,
            ) in _WORD_WHITESPACE_CASES.items()
        ]
    lines = [
        _whitespace_line(
            f"{kind}{case}",
            content,
            space_before,
            space_after,
            lambda inner, tags=tags: tags[0] + inner + tags[1],
        )
        for kind, tags in _HTML_INLINE.items()
        for case, (content, space_before, space_after) in _WHITESPACE_CASES.items()
    ]
    lines.append(
        WhitespaceRunLine(
            "autolinkprior",
            f'<a href="{WHITESPACE_RUN_HREF}"> {WHITESPACE_RUN_HREF} </a>',
            "autolinknext",
            ("autolinkprior", WHITESPACE_RUN_HREF, "autolinknext"),
        )
    )
    return lines


def _word_whitespace_run(kind: str, content: str) -> str:
    if kind == "link":
        return f'<w:hyperlink r:id="rId2"><w:r>{content}</w:r></w:hyperlink>'
    return f"<w:r><w:rPr>{_WORD_RUN_PROPERTIES[kind]}</w:rPr>{content}</w:r>"


def write_whitespace_run_report(path: Path) -> None:
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"'
        f' Target="{EDGE_SPACE_HREF}" TargetMode="External"/></Relationships>\n'
    )

    body = _docx_paragraph(SENTINELS["whitespace_run_report.docx"]) + "".join(
        f"<w:p>{_word_run(line.before)}{line.run}{_word_run(line.after)}</w:p>"
        for line in whitespace_run_lines("whitespace_run_report.docx")
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", _DOCX_CONTENT_TYPES),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
        ],
    )


def _whitespace_run_html(name: str) -> str:
    return "\n".join(
        f"  <p>{line.before}{line.run}{line.after}</p>"
        for line in whitespace_run_lines(name)
    )


def write_whitespace_run_book(path: Path) -> None:
    opf = _EPUB_OPF.replace("webshot-office-fixture", "webshot-whitespace-run-fixture")
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter One</title></head>
<body>
  <p>{SENTINELS["whitespace_run_book.epub"]}</p>
{_whitespace_run_html("whitespace_run_book.epub")}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", opf),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


def write_whitespace_run_archive(path: Path) -> None:
    """An HTML page in a zip, as `write_edge_space_archive` explains."""
    page = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Whitespace runs</title></head>
<body>
  <p>{SENTINELS["whitespace_run_archive.zip"]}</p>
{_whitespace_run_html("whitespace_run_archive.zip")}
</body></html>
"""
    _write_archive(path, [("pages/whitespace-runs.html", page)])


# --------------------------------------------------------------------------- #
# Quotations and definition lists side by side (docs/09 P63)
# --------------------------------------------------------------------------- #


def _block_html(block: tuple[str, list[Any]]) -> str:
    """A quotation's or a definition list's HTML: `("blockquote", words)`
    holds a paragraph of each word, `("dl", pairs)` a term and a definition
    for each."""
    tag, content = block
    if tag == "blockquote":
        return (
            "<blockquote>" + "".join(f"<p>{w}</p>" for w in content) + "</blockquote>"
        )
    return "<dl>" + "".join(f"<dt>{t}</dt><dd>{d}</dd>" for t, d in content) + "</dl>"


#: Each pair, by what stands between or around it in the chapter, and the
#: blocks it holds.  Python-Markdown read every pair but the last as one block
#: until the bridge kept them apart; the last, a quotation and a definition
#: list, it never joined.  A pair in a quotation or a definition is written
#: inside one.  The words are made up.
ADJACENT_BLOCKS: dict[str, tuple[str, list[tuple[str, list[Any]]]]] = {
    "Two quotations": (
        "{0}{1}",
        [("blockquote", ["Arvolt", "Belsk"]), ("blockquote", ["Corwine"])],
    ),
    "Three quotations": (
        "{0}{1}{2}",
        [
            ("blockquote", ["Dremmet"]),
            ("blockquote", ["Elvast"]),
            ("blockquote", ["Frindle"]),
        ],
    ),
    "Two definition lists": (
        "{0}{1}",
        [
            ("dl", [("Gantry", "Holvet"), ("Iskel", "Jarrow")]),
            ("dl", [("Kelp", "Lunt")]),
        ],
    ),
    "Nothing that shows between them": (
        "{0}\n  <div></div><p></p><!-- a note --><br/>{1}",
        [("dl", [("Morrow", "Nesk")]), ("dl", [("Orlin", "Pask")])],
    ),
    "Each in a division of its own": (
        "<div>{0}</div><div>{1}</div>",
        [("blockquote", ["Quillet"]), ("blockquote", ["Rusket"])],
    ),
    "Inside one quotation": (
        "<blockquote>{0}{1}</blockquote>",
        [("blockquote", ["Sallow"]), ("blockquote", ["Tamsk"])],
    ),
    "Inside one definition": (
        "<dl><dt>Umbrel</dt><dd>{0}{1}</dd></dl>",
        [("dl", [("Vintry", "Wexel")]), ("dl", [("Yarrel", "Zoffet")])],
    ),
    "A quotation, then a definition list": (
        "{0}{1}",
        [("blockquote", ["Abbot"]), ("dl", [("Brisk", "Cawdle")])],
    ),
}

#: What chapter one ends with and chapter two begins with.
ADJACENT_BLOCKS_ACROSS: list[tuple[str, list[Any]]] = [
    ("blockquote", ["Dunnage"]),
    ("blockquote", ["Effrin"]),
]


def adjacent_blocks_page(
    sections: dict[str, tuple[str, list[tuple[str, list[Any]]]]],
) -> list[tuple[str, list[str]]]:
    """The quotations and definition lists a page of `sections` holds, in the
    order their tags open: each tag and its words, a term and its definition
    one after the other."""
    page: list[tuple[str, list[str]]] = []
    for layout, blocks in sections.values():
        if layout.startswith("<blockquote>"):
            page.append(("blockquote", []))
        if layout.startswith("<dl>"):
            page.append(("dl", ["Umbrel"]))
        for tag, content in blocks:
            words = (
                content
                if tag == "blockquote"
                else [w for pair in content for w in pair]
            )
            page.append((tag, list(words)))
    return page


#: The book's quotations and definition lists as its page holds them.
ADJACENT_BLOCKS_PAGE = adjacent_blocks_page(ADJACENT_BLOCKS) + [
    (tag, list(words)) for tag, words in ADJACENT_BLOCKS_ACROSS
]


def _adjacent_blocks_sections() -> str:
    return "".join(
        f"<h2>{heading}</h2>" + layout.format(*map(_block_html, blocks))
        for heading, (layout, blocks) in ADJACENT_BLOCKS.items()
    )


def write_adjacent_blocks_book(path: Path) -> None:
    """Pairs of quotations and of definition lists, in two chapters.

    The first chapter ends with a quotation and the second begins with one:
    MarkItDown converts each chapter on its own and joins them with a blank
    line, so the two met only once the book was whole.
    """
    first = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Blocks</title></head>
<body>
  <h1>Quotations Side by Side</h1>
  <p>{SENTINELS["adjacent_blocks_book.epub"]}</p>
  {_adjacent_blocks_sections()}
  <h2>Across two chapters</h2>
  {_block_html(ADJACENT_BLOCKS_ACROSS[0])}
</body></html>
"""
    second = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>More</title></head>
<body>
  {_block_html(ADJACENT_BLOCKS_ACROSS[1])}
  <p>The second chapter goes on after its quotation.</p>
</body></html>
"""
    opf = """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="book-id">urn:uuid:webshot-adjacent-blocks-fixture</dc:identifier>
    <dc:title>Quotations Side by Side</dc:title>
    <dc:language>en</dc:language>
  </metadata>
  <manifest>
    <item id="one" href="one.xhtml" media-type="application/xhtml+xml"/>
    <item id="two" href="two.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="one"/><itemref idref="two"/></spine>
</package>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", opf),
            ("OEBPS/one.xhtml", first),
            ("OEBPS/two.xhtml", second),
        ],
        first_stored=True,
    )


#: mammoth writes a quotation or a definition list only where a style map says
#: to.  A Word file can carry one, and MarkItDown reads it: every `Quote`
#: paragraph a quotation of its own, and a `NewTerm` beginning a definition
#: list where a `Term` goes on with the one before.
ADJACENT_BLOCKS_STYLE_MAP = (
    "p.Quote => blockquote:fresh > p:fresh\n"
    "p.NewTerm => dl:fresh > dt:fresh\n"
    "p.Term => dl > dt:fresh\n"
    "p.Definition => dl > dd:fresh\n"
)

#: The report's paragraphs, each its style and its text; a style of None is
#: body text.
ADJACENT_BLOCKS_REPORT_PARAGRAPHS: list[tuple[str | None, str]] = [
    ("Quote", "Fallow"),
    ("Quote", "Gimbrel"),
    (None, "Two definition lists follow."),
    ("NewTerm", "Hessock"),
    ("Definition", "Ilvane"),
    ("Term", "Jessamy"),
    ("Definition", "Kildrum"),
    ("NewTerm", "Lorrimer"),
    ("Definition", "Mossel"),
]

ADJACENT_BLOCKS_REPORT_PAGE: list[tuple[str, list[str]]] = [
    ("blockquote", ["Fallow"]),
    ("blockquote", ["Gimbrel"]),
    ("dl", ["Hessock", "Ilvane", "Jessamy", "Kildrum"]),
    ("dl", ["Lorrimer", "Mossel"]),
]


def write_adjacent_blocks_report(path: Path) -> None:
    """A Word file whose embedded style map writes quotations and definition
    lists side by side.

    Laid out as mammoth's own `embed_style_map` writes one: the map at
    `mammoth/style-map`, named in the content types and the document's
    relationships.
    """
    body = (
        _docx_paragraph("Quotations Side by Side", "Heading1")
        + _docx_paragraph(SENTINELS["adjacent_blocks_report.docx"])
        + "".join(
            _docx_paragraph(text, style)
            for style, text in ADJACENT_BLOCKS_REPORT_PARAGRAPHS
        )
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    content_types = _DOCX_CONTENT_TYPES.replace(
        "</Types>",
        '  <Override PartName="/mammoth/style-map"'
        ' ContentType="text/prs.mammoth.style-map"/>\n</Types>',
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rMammothStyleMap" Type="http://schemas.zwobble.org/mammoth/style-map"'
        ' Target="/mammoth/style-map"/></Relationships>\n'
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", content_types),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
            ("mammoth/style-map", ADJACENT_BLOCKS_STYLE_MAP),
        ],
    )


#: An HTML page in the archive: a zip member is read by the same writer.
_ADJACENT_BLOCKS_MEMBER_PAGE = f"""<!DOCTYPE html>
<html><head><title>Side by side</title></head>
<body>
  <p>{SENTINELS["adjacent_blocks_archive.zip"]}</p>
  <blockquote><p>Nettle</p></blockquote><blockquote><p>Orpin</p></blockquote>
  <dl><dt>Pelter</dt><dd>Quist</dd></dl><dl><dt>Rennet</dt><dd>Sorrel</dd></dl>
</body></html>
"""

#: What the archive's page holds: its own page's, then the report's.
ADJACENT_BLOCKS_ARCHIVE_PAGE: list[tuple[str, list[str]]] = [
    ("blockquote", ["Nettle"]),
    ("blockquote", ["Orpin"]),
    ("dl", ["Pelter", "Quist"]),
    ("dl", ["Rennet", "Sorrel"]),
    *ADJACENT_BLOCKS_REPORT_PAGE,
]


def write_adjacent_blocks_archive(path: Path) -> None:
    """An HTML page of quotations and definition lists side by side, and the
    report as it is."""
    with tempfile.TemporaryDirectory() as scratch:
        report = Path(scratch) / "adjacent_blocks_report.docx"
        write_adjacent_blocks_report(report)
        _write_archive(
            path,
            [
                ("notes/side-by-side.html", _ADJACENT_BLOCKS_MEMBER_PAGE),
                ("reports/adjacent_blocks_report.docx", report.read_bytes()),
            ],
        )


# --------------------------------------------------------------------------- #
# Word list numbers (docs/09 P31)
# --------------------------------------------------------------------------- #


def _numbering_level(level: int, fmt: str = "decimal", start: int | None = 1) -> str:
    start_xml = "" if start is None else f'<w:start w:val="{start}"/>'
    return (
        f'<w:lvl w:ilvl="{level}">{start_xml}<w:numFmt w:val="{fmt}"/>'
        f'<w:lvlText w:val="%{level + 1}."/><w:lvlJc w:val="left"/></w:lvl>'
    )


#: Each abstract list definition: its first level's format and start.  `None`
#: is a level with no `w:start`, which Word counts from 0.
_NUMBERED_DEFINITIONS = {
    1: ("decimal", 3),
    2: ("decimal", 1),
    3: ("decimal", 1),
    4: ("decimal", 1),
    5: ("decimal", 1),
    6: ("decimal", 1),
    7: ("decimal", None),
    8: ("decimal", 1),
    9: ("lowerLetter", 2),
    10: ("decimal", 5),
    11: ("decimal", 1),
    12: ("none", 1),
}
#: Each list instance (`w:num`): its definition and, where it has one, the
#: `w:startOverride` of its first level.
_NUMBERED_INSTANCES = {
    1: (1, None),
    2: (2, 4),
    3: (3, None),
    4: (4, None),
    5: (4, None),
    6: (5, None),
    7: (6, None),
    8: (7, None),
    9: (8, None),
    10: (9, None),
    11: (10, None),
    12: (11, None),
    13: (12, None),
}
#: The document, a section a heading: each paragraph's text and the list
#: instance numbering it, if any.  Word's numbers for each are in
#: `NUMBERED_PAGE`, measured with Word 16.113 (docs/09 P31).
NUMBERED_SECTIONS: list[tuple[str, list[tuple[str, int | None]]]] = [
    ("A list that starts at three", [("Brindle", 1), ("Carrow", 1)]),
    ("A list restarted at four", [("Dorvent", 2), ("Elkish", 2)]),
    (
        "A list that goes on past a paragraph",
        [("Fennow", 3), ("Gaskel", 3), ("EQUATION", None), ("Hollin", 3)],
    ),
    (
        "Two lists of one definition",
        [("Istrel", 4), ("Jorvik", 4), ("Between them, a note.", None), ("Kelmar", 5)],
    ),
    ("Two lists that meet", [("Lunt", 6), ("Morrow", 6), ("Nebbit", 7), ("Orlen", 7)]),
    ("A list counted from zero", [("Pellam", 8), ("Quist", 8)]),
    ("A list with an empty item", [("Rusk", 9), ("", 9), ("Sorrel", 9)]),
    ("A lettered list", [("Tamsin", 10), ("Ulbrey", 10)]),
    ("A list its style numbers", [("Vellum", -1), ("Wexford", -1)]),
    ("A list Word shows no numbers on", [("Yarrow", 13), ("Zennor", 13)]),
]
#: The heading numbered as a list, which the page shows without its number.
NUMBERED_HEADING = ("Brackwater", 12)
#: content.md's lists, each item as `N. text`, with Word's numbers: a lettered
#: item by its ordinal, and the list Word shows no numbers on as the page
#: numbers it, which the manifest says.
NUMBERED_PAGE = [
    ["3. Brindle", "4. Carrow"],
    ["4. Dorvent", "5. Elkish"],
    ["1. Fennow", "2. Gaskel"],
    ["3. Hollin"],
    ["1. Istrel", "2. Jorvik"],
    ["3. Kelmar"],
    ["1. Lunt", "2. Morrow"],
    ["1. Nebbit", "2. Orlen"],
    ["0. Pellam", "1. Quist"],
    ["1. Rusk"],
    ["3. Sorrel"],
    ["2. Tamsin", "3. Ulbrey"],
    ["5. Vellum", "6. Wexford"],
    ["1. Yarrow", "2. Zennor"],
]


def write_numbered_report(path: Path) -> None:
    """Word lists whose numbers mammoth does not carry.

    A list starting past 1, by its definition and by an override; one going on
    past a paragraph, one of whose paragraphs holds an equation so that the
    document is rewritten before it is numbered (docs/09 P27); two instances
    of one definition, which Word counts as one list; two lists that meet; a
    definition with no start; an empty item; a lettered list; a list numbered
    by its paragraph style; a heading numbered as a list; and a list whose
    format is `none`.  Every word is made up.
    """
    abstracts = "".join(
        f'<w:abstractNum w:abstractNumId="{number}">'
        + (
            '<w:lvl w:ilvl="0"><w:start w:val="5"/><w:numFmt w:val="decimal"/>'
            '<w:pStyle w:val="ListNumber"/><w:lvlText w:val="%1."/></w:lvl>'
            if number == 10
            else _numbering_level(0, fmt, start)
        )
        + "</w:abstractNum>"
        for number, (fmt, start) in _NUMBERED_DEFINITIONS.items()
    )
    instances = "".join(
        f'<w:num w:numId="{number}"><w:abstractNumId w:val="{definition}"/>'
        + (
            ""
            if override is None
            else f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="{override}"/></w:lvlOverride>'
        )
        + "</w:num>"
        for number, (definition, override) in _NUMBERED_INSTANCES.items()
    )
    numbering = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<w:numbering xmlns:w="{_OOXML["w"]}">{abstracts}{instances}</w:numbering>\n'
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<w:styles xmlns:w="{_OOXML["w"]}">'
        '<w:style w:type="paragraph" w:default="1" w:styleId="Normal">'
        '<w:name w:val="Normal"/></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/>'
        '<w:basedOn w:val="Normal"/></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/>'
        '<w:basedOn w:val="Normal"/></w:style>'
        '<w:style w:type="paragraph" w:styleId="ListNumber"><w:name w:val="List Number"/>'
        '<w:basedOn w:val="Normal"/><w:pPr><w:numPr><w:numId w:val="11"/></w:numPr>'
        "</w:pPr></w:style></w:styles>\n"
    )

    def paragraph(text: str, instance: int | None, style: str | None = None) -> str:
        properties = f'<w:pStyle w:val="{style}"/>' if style else ""
        if instance is not None and instance > 0:
            properties += (
                f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{instance}"/></w:numPr>'
            )
        if instance == -1:
            properties += '<w:pStyle w:val="ListNumber"/>'
        if text == "EQUATION":
            content = (
                _word_run("Between the steps, ")
                + _math(_word_math_run("x") + _word_math_run("≠") + _word_math_run("y"))
                + _word_run(".")
            )
        else:
            content = _word_run(text) if text else ""
        return f"<w:p><w:pPr>{properties}</w:pPr>{content}</w:p>"

    heading, instance = NUMBERED_HEADING
    body = (
        paragraph("Numbered Steps", None, "Heading1")
        + paragraph(SENTINELS["numbered_report.docx"], None)
        + paragraph(heading, instance, "Heading1")
        + "".join(
            paragraph(title, None, "Heading2")
            + "".join(paragraph(text, item) for text, item in items)
            for title, items in NUMBERED_SECTIONS
        )
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering"'
        ' Target="numbering.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"'
        ' Target="styles.xml"/></Relationships>\n'
    )
    content_types = _DOCX_CONTENT_TYPES.replace(
        "</Types>",
        '  <Override PartName="/word/numbering.xml" ContentType="application/'
        'vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>\n'
        '  <Override PartName="/word/styles.xml" ContentType="application/'
        'vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>\n'
        "</Types>",
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", content_types),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
            ("word/numbering.xml", numbering),
            ("word/styles.xml", styles),
        ],
    )


# --------------------------------------------------------------------------- #
# Terms in a row, and definitions after one of more than one block (docs/09 P64)
# --------------------------------------------------------------------------- #

#: A definition list as the page should hold it: each term and definition in
#: order, by its tag and its text, a nested list's text with the definition
#: it sits in.  Lists in the order a page has them, an outer before its inner.
DefinitionList = list[tuple[str, str]]

#: The book's definition lists, by what each section shows.  markdownify
#: writes a blank line before every term and none before a definition, so in
#: each of these but the last a term became a paragraph, or a definition text
#: inside the one before it.  The words are made up.
DEFINITIONS_BOOK_SECTIONS = {
    "Two terms for one definition": (
        "<dl><dt>Arvesh</dt><dt>Bolwick</dt><dd>Cendral</dd></dl>"
    ),
    "Three terms for a definition of two paragraphs": (
        "<dl><dt>Dovrane</dt><dt>Elmsark</dt><dt>Fenwold</dt>"
        "<dd><p>Garnick</p><p>Holvane</p></dd></dl>"
    ),
    "Two definitions of two paragraphs": (
        "<dl><dt>Istrel</dt><dd><p>Jorvane</p><p>Kelmond</p></dd>"
        "<dd><p>Lindrow</p><p>Morvish</p></dd></dl>"
    ),
    "A definition after one holding a list": (
        "<dl><dt>Norbel</dt><dd><p>Ostrin</p><ul><li>Pelvane</li></ul></dd>"
        "<dd>Quinlow</dd></dl>"
    ),
    "Two terms after a definition of two paragraphs": (
        "<dl><dt>Rastel</dt><dd><p>Sorvin</p><p>Telmark</p></dd>"
        "<dt>Ulvane</dt><dt>Varnow</dt><dd>Wendrel</dd></dl>"
    ),
    "Terms in a quotation": (
        "<blockquote><dl><dt>Yarrel</dt><dt>Zelkin</dt><dd>Ambrel</dd></dl>"
        "</blockquote>"
    ),
    "A definition list in a definition": (
        "<dl><dt>Brisvel</dt><dd><p>Corvane</p><dl><dt>Dellow</dt>"
        "<dt>Emberwick</dt><dd>Falstan</dd></dl></dd><dd>Gorwin</dd></dl>"
    ),
    "Terms in a list item": (
        "<ul><li>Harnel<dl><dt>Ivorne</dt><dt>Jessep</dt><dd>Kettrel</dd></dl>"
        "</li></ul>"
    ),
    "A term on its own, as before": (
        "<dl><dt>Lumbrel</dt><dd>Mavish</dd><dt>Nolwen</dt><dd>Orvane</dd></dl>"
    ),
}

#: `DEFINITIONS_BOOK_SECTIONS`' definition lists as the page holds them.
DEFINITIONS_BOOK_PAGE: list[DefinitionList] = [
    [("dt", "Arvesh"), ("dt", "Bolwick"), ("dd", "Cendral")],
    [
        ("dt", "Dovrane"),
        ("dt", "Elmsark"),
        ("dt", "Fenwold"),
        ("dd", "Garnick Holvane"),
    ],
    [("dt", "Istrel"), ("dd", "Jorvane Kelmond"), ("dd", "Lindrow Morvish")],
    [("dt", "Norbel"), ("dd", "Ostrin Pelvane"), ("dd", "Quinlow")],
    [
        ("dt", "Rastel"),
        ("dd", "Sorvin Telmark"),
        ("dt", "Ulvane"),
        ("dt", "Varnow"),
        ("dd", "Wendrel"),
    ],
    [("dt", "Yarrel"), ("dt", "Zelkin"), ("dd", "Ambrel")],
    [
        ("dt", "Brisvel"),
        ("dd", "Corvane Dellow Emberwick Falstan"),
        ("dd", "Gorwin"),
    ],
    [("dt", "Dellow"), ("dt", "Emberwick"), ("dd", "Falstan")],
    [("dt", "Ivorne"), ("dt", "Jessep"), ("dd", "Kettrel")],
    [("dt", "Lumbrel"), ("dd", "Mavish"), ("dt", "Nolwen"), ("dd", "Orvane")],
]


def write_definitions_book(path: Path) -> None:
    sections = "".join(
        f"<h2>{heading}</h2>{lists}"
        for heading, lists in DEFINITIONS_BOOK_SECTIONS.items()
    )
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Terms</title></head>
<body>
  <h1>Terms and Their Definitions</h1>
  <p>{SENTINELS["definitions_book.epub"]}</p>
  {sections}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


#: An HTML page in the archive: a zip member is read by the same writer.
_DEFINITIONS_MEMBER_PAGE = f"""<!DOCTYPE html>
<html><head><title>Glossary</title></head>
<body>
  <p>{SENTINELS["definitions_archive.zip"]}</p>
  <dl>
    <dt>Pardrel</dt>
    <dt>Quessick</dt>
    <dd><p>Rombert</p><p>Sallow</p></dd>
    <dd><p>Tamwick</p><p>Undrel</p></dd>
  </dl>
</body></html>
"""

#: What the archive's page holds: its own page's list, then the book's.
DEFINITIONS_ARCHIVE_PAGE: list[DefinitionList] = [
    [
        ("dt", "Pardrel"),
        ("dt", "Quessick"),
        ("dd", "Rombert Sallow"),
        ("dd", "Tamwick Undrel"),
    ],
    *DEFINITIONS_BOOK_PAGE,
]


def write_definitions_archive(path: Path) -> None:
    """An HTML page with a definition list, and the definitions book as it is."""
    with tempfile.TemporaryDirectory() as scratch:
        book = Path(scratch) / "definitions_book.epub"
        write_definitions_book(book)
        _write_archive(
            path,
            [
                ("notes/glossary.html", _DEFINITIONS_MEMBER_PAGE),
                ("books/definitions_book.epub", book.read_bytes()),
            ],
        )


# --------------------------------------------------------------------------- #
# Lists in a quotation in a list item (docs/09 P74)
# --------------------------------------------------------------------------- #

#: The book's lists, by what each section shows.  markdownify wrote a list in a
#: quotation in an item as the item's own, with no blank line after it, so
#: every section but the last lost its shape on the page.  The last, with no
#: item around its quotation, was always right.  The words are made up.
QUOTED_LISTS_SECTIONS = {
    "Text after the list": (
        "<ul><li>Astrel<blockquote><ul><li>Brimmock</li></ul>"
        "and Caddrel after it</blockquote></li><li>Dovrel</li></ul>"
    ),
    "A bulleted list, then a numbered one": (
        "<ul><li>Ellwick<blockquote><ul><li>Fennow</li></ul>"
        "<ol><li>Gorrin</li></ol></blockquote></li></ul>"
    ),
    "Two numbered lists": (
        "<ol><li>Hastrin<blockquote><ol><li>Ilvane</li></ol>"
        "<ol><li>Jessock</li></ol></blockquote></li><li>Kendrow</li></ol>"
    ),
    "Two bulleted lists": (
        "<ul><li>Ludrel<blockquote><ul><li>Mosket</li></ul>"
        "<ul><li>Nestrow</li></ul></blockquote></li></ul>"
    ),
    "Under a number of two digits": (
        '<ol start="10"><li>Fendrel<blockquote><ol><li>Garvish</li></ol>'
        "<ol><li>Holdrin</li></ol></blockquote></li></ol>"
    ),
    "In an item two lists deep": (
        "<ul><li>Orvel<ol><li>Pandrish<blockquote><ol><li>Quessel</li></ol>"
        "and Rostwick after it</blockquote></li></ol></li></ul>"
    ),
    "In an item three lists deep": (
        "<ol><li>Sarvel<ul><li>Tobrin<ol><li>Ulvesk<blockquote>"
        "<ul><li>Vandrel</li></ul><ol><li>Wessock</li></ol>"
        "</blockquote></li></ol></li></ul></li></ol>"
    ),
    "A quotation in a quotation": (
        "<ul><li>Yarrel<blockquote><blockquote><ul><li>Zendish</li></ul>"
        "and Abrovel after it</blockquote></blockquote></li></ul>"
    ),
    "A list nested in the quotation's list": (
        "<ul><li>Bellock<blockquote><ul><li>Corvash<ol><li>Delvish</li></ol>"
        "</li></ul>and Erdwick after it</blockquote></li></ul>"
    ),
    "No item around the quotation": (
        "<blockquote><ul><li>Istow</li></ul>and Jarvel after it</blockquote>"
    ),
}


def _quoted_lists_body(sentinel: str) -> str:
    sections = "".join(
        f"<h2>{heading}</h2>{lists}" for heading, lists in QUOTED_LISTS_SECTIONS.items()
    )
    return f"<h1>Lists Within Quotations</h1><p>{sentinel}</p>{sections}"


def write_quoted_lists_book(path: Path) -> None:
    body = _quoted_lists_body(SENTINELS["quoted_lists_book.epub"])
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Quoted</title></head>
<body>{body}</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", _EPUB_OPF),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


def write_quoted_lists_archive(path: Path) -> None:
    """The book's chapter as an HTML page: a zip member is read by the same
    writer, without the book's reader around it."""
    body = _quoted_lists_body(SENTINELS["quoted_lists_archive.zip"])
    page = (
        "<!DOCTYPE html>\n<html><head><title>Quoted</title></head>"
        f"<body>{body}</body></html>\n"
    )
    _write_archive(path, [("notes/quoted.html", page)])


# --------------------------------------------------------------------------- #
# Line breaks — a `<br>`, or Word's `w:br`, in each place text flows (P77)
# --------------------------------------------------------------------------- #

#: The made-up endings of a case's words: its label, then one of these.
_LINE_BREAK_ENDINGS = ("brom", "drel", "fisk", "gorv", "hult")
#: Where the Word fixture's link points; `.invalid` never resolves.
LINE_BREAK_HREF = "https://example.invalid/break"

#: Each case as HTML, for a book or a page: its markup, with `{0}` and on for
#: its words, and the lines the page shows, top to bottom.  A line the
#: document leaves empty is not listed: what is checked is which words share
#: a line and which a line break parts.  Two `<br>`s in a paragraph are the
#: paragraph break they were; `runbold`'s is one docling drops otherwise.
_LINE_BREAK_HTML = {
    "para": ("<p>{0}<br/>{1}</p>", ("{0}", "{1}")),
    "pretty": ("<p>{0}<br/>\n    {1}</p>", ("{0}", "{1}")),
    "double": ("<p>{0}<br/><br/>{1}</p>", ("{0}", "{1}")),
    "runbold": ("<p>{0}<br/><br/><strong>{1}</strong> {2}</p>", ("{0}", "{1} {2}")),
    "lead": ("<p><br/>{0} {1}</p>", ("{0} {1}",)),
    "trail": ("<p>{0} {1}<br/></p>", ("{0} {1}",)),
    "only": ("<p>{0}</p><p><br/></p><p>{1}</p>", ("{0}", "{1}")),
    "bold": ("<p>{0} <strong>{1}<br/>{2}</strong> {3}</p>", ("{0} {1}", "{2} {3}")),
    "boldrun": (
        "<p>{0} <strong>{1}<br/><br/>{2}</strong> {3}</p>",
        ("{0} {1}", "{2} {3}"),
    ),
    "boldtrail": ("<p>{0} <strong>{1}<br/></strong>{2}</p>", ("{0} {1}", "{2}")),
    "boldlead": ("<p>{0}<strong><br/>{1}</strong> {2}</p>", ("{0}", "{1} {2}")),
    "boldalone": ("<p>{0}<strong><br/></strong>{1}</p>", ("{0}", "{1}")),
    "italic": ("<p>{0} <em>{1}<br/>{2}</em> {3}</p>", ("{0} {1}", "{2} {3}")),
    "link": (
        f'<p>{{0}} <a href="{LINE_BREAK_HREF}">{{1}}<br/>{{2}}</a> {{3}}</p>',
        ("{0} {1}", "{2} {3}"),
    ),
    "item": (
        "<ul><li>{0}<br/>{1}</li><li>{2}<br/><br/>{3}</li><li>{4}</li></ul>",
        ("{0}", "{1}", "{2}", "{3}", "{4}"),
    ),
    "numbered": ("<ol><li>{0}<br/>{1}</li><li>{2}</li></ol>", ("{0}", "{1}", "{2}")),
    "quote": ("<blockquote><p>{0}<br/>{1}</p></blockquote>", ("{0}", "{1}")),
    "cell": (
        "<table><tr><th>{0}<br/>{1}</th></tr><tr><td>{2}<br/>{3}</td></tr></table>",
        ("{0}", "{1}", "{2}", "{3}"),
    ),
    "openers": (
        "<p>{0}<br/>- {1}<br/>3. {2}<br/># {3}</p>",
        ("{0}", "- {1}", "3. {2}", "# {3}"),
    ),
    # A heading's stays a space, which the PDF's bookmarks would read run
    # together; a `<pre>`'s a newline, which the capture's print CSS draws;
    # and inline code's a space, as every line break in inline code is, which
    # is written on one line so that none of its lines can open a block
    # (docs/09 P59).
    "heading": ("<h2>{0}<br/>{1}</h2>", ("{0} {1}",)),
    "code": ("<p>{0} <code>{1}<br/>{2}</code> {3}</p>", ("{0} {1} {2} {3}",)),
    "block": ("<pre>{0}<br/>{1}</pre>", ("{0}", "{1}")),
}


def _word_text(text: str) -> str:
    return f'<w:t xml:space="preserve">{text}</w:t>'


def _word_formatted(properties: str, content: str) -> str:
    return f"<w:r><w:rPr>{properties}</w:rPr>{content}</w:r>"


_WORD_BREAK = "<w:br/>"

#: The same, as Word holds them: `w:br` in a run, which mammoth writes as
#: `<br />`.  Word has no list without a numbering part, and no quotation.
_LINE_BREAK_WORD = {
    "para": (
        f"<w:p>{_word_run('{0}')}<w:r>{_WORD_BREAK}</w:r>{_word_run('{1}')}</w:p>",
        ("{0}", "{1}"),
    ),
    "inrun": (
        f"<w:p><w:r>{_word_text('{0}')}{_WORD_BREAK}{_word_text('{1}')}</w:r></w:p>",
        ("{0}", "{1}"),
    ),
    "double": (
        f"<w:p><w:r>{_word_text('{0}')}{_WORD_BREAK * 2}{_word_text('{1}')}</w:r></w:p>",
        ("{0}", "{1}"),
    ),
    "runbold": (
        f"<w:p><w:r>{_word_text('{0}')}{_WORD_BREAK * 2}</w:r>"
        + _word_formatted("<w:b/>", _word_text("{1}"))
        + f"{_word_run(' {2}')}</w:p>",
        ("{0}", "{1} {2}"),
    ),
    "lead": (
        f"<w:p><w:r>{_WORD_BREAK}{_word_text('{0} {1}')}</w:r></w:p>",
        ("{0} {1}",),
    ),
    "trail": (
        f"<w:p><w:r>{_word_text('{0} {1}')}{_WORD_BREAK}</w:r></w:p>",
        ("{0} {1}",),
    ),
    "only": (
        f"<w:p>{_word_run('{0}')}</w:p><w:p><w:r>{_WORD_BREAK}</w:r></w:p>"
        f"<w:p>{_word_run('{1}')}</w:p>",
        ("{0}", "{1}"),
    ),
    "bold": (
        f"<w:p>{_word_run('{0} ')}"
        + _word_formatted("<w:b/>", _word_text("{1}") + _WORD_BREAK + _word_text("{2}"))
        + f"{_word_run(' {3}')}</w:p>",
        ("{0} {1}", "{2} {3}"),
    ),
    "boldrun": (
        f"<w:p>{_word_run('{0} ')}"
        + _word_formatted(
            "<w:b/>", _word_text("{1}") + _WORD_BREAK * 2 + _word_text("{2}")
        )
        + f"{_word_run(' {3}')}</w:p>",
        ("{0} {1}", "{2} {3}"),
    ),
    "boldtrail": (
        f"<w:p>{_word_run('{0} ')}"
        + _word_formatted("<w:b/>", _word_text("{1}") + _WORD_BREAK)
        + f"{_word_run('{2}')}</w:p>",
        ("{0} {1}", "{2}"),
    ),
    "boldlead": (
        f"<w:p>{_word_run('{0}')}"
        + _word_formatted("<w:b/>", _WORD_BREAK + _word_text("{1}"))
        + f"{_word_run(' {2}')}</w:p>",
        ("{0}", "{1} {2}"),
    ),
    "boldalone": (
        f"<w:p>{_word_run('{0}')}"
        + _word_formatted("<w:b/>", _WORD_BREAK)
        + f"{_word_run('{1}')}</w:p>",
        ("{0}", "{1}"),
    ),
    "italic": (
        f"<w:p>{_word_run('{0} ')}"
        + _word_formatted("<w:i/>", _word_text("{1}") + _WORD_BREAK + _word_text("{2}"))
        + f"{_word_run(' {3}')}</w:p>",
        ("{0} {1}", "{2} {3}"),
    ),
    "link": (
        f"<w:p>{_word_run('{0} ')}"
        '<w:hyperlink r:id="rId2"><w:r>'
        + _word_text("{1}")
        + _WORD_BREAK
        + _word_text("{2}")
        + f"</w:r></w:hyperlink>{_word_run(' {3}')}</w:p>",
        ("{0} {1}", "{2} {3}"),
    ),
    "cell": (
        "<w:tbl><w:tblPr/>"
        + "".join(
            f"<w:tr><w:tc><w:tcPr/><w:p><w:r>{_word_text(first)}{_WORD_BREAK}"
            f"{_word_text(second)}</w:r></w:p></w:tc></w:tr>"
            for first, second in (("{0}", "{1}"), ("{2}", "{3}"))
        )
        + "</w:tbl>",
        ("{0}", "{1}", "{2}", "{3}"),
    ),
    "openers": (
        "<w:p><w:r>"
        + _WORD_BREAK.join(
            _word_text(text) for text in ("{0}", "- {1}", "3. {2}", "# {3}")
        )
        + "</w:r></w:p>",
        ("{0}", "- {1}", "3. {2}", "# {3}"),
    ),
    "heading": (
        '<w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r>'
        f"{_word_text('{0}')}{_WORD_BREAK}{_word_text('{1}')}</w:r></w:p>",
        ("{0} {1}",),
    ),
}


@dataclass(frozen=True)
class LineBreakCase:
    """One case of a `line_break_*` fixture: markup with line breaks in it,
    and the lines the page shows, each as its words."""

    label: str
    markup: str
    lines: tuple[tuple[str, ...], ...]

    @property
    def words(self) -> tuple[str, ...]:
        """The case's made-up words, in order: what content.md must keep apart."""
        return tuple(
            word for line in self.lines for word in line if word.startswith(self.label)
        )


def line_break_cases(name: str) -> list[LineBreakCase]:
    """The cases of the `line_break_*` fixture called `name`."""
    table = _LINE_BREAK_WORD if name.endswith(".docx") else _LINE_BREAK_HTML
    cases = []
    for label, (markup, lines) in table.items():
        words = [f"{label}{ending}" for ending in _LINE_BREAK_ENDINGS]
        cases.append(
            LineBreakCase(
                label,
                markup.format(*words),
                tuple(tuple(line.format(*words).split()) for line in lines),
            )
        )
    return cases


def write_line_break_report(path: Path) -> None:
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"'
        f' Target="{LINE_BREAK_HREF}" TargetMode="External"/></Relationships>\n'
    )
    body = _docx_paragraph(SENTINELS["line_break_report.docx"]) + "".join(
        case.markup for case in line_break_cases("line_break_report.docx")
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", _DOCX_CONTENT_TYPES),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
        ],
    )


def _line_break_html(name: str) -> str:
    return "\n".join(f"  {case.markup}" for case in line_break_cases(name))


def write_line_break_book(path: Path) -> None:
    opf = _EPUB_OPF.replace("webshot-office-fixture", "webshot-line-break-fixture")
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter One</title></head>
<body>
  <p>{SENTINELS["line_break_book.epub"]}</p>
{_line_break_html("line_break_book.epub")}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", opf),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


def write_line_break_archive(path: Path) -> None:
    """An HTML page in a zip, as `write_edge_space_archive` explains."""
    page = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Line breaks</title></head>
<body>
  <p>{SENTINELS["line_break_archive.zip"]}</p>
{_line_break_html("line_break_archive.zip")}
</body></html>
"""
    _write_archive(path, [("pages/line-breaks.html", page)])


# --------------------------------------------------------------------------- #
# Bare links — a link whose text is its target (docs/09 P75)
# --------------------------------------------------------------------------- #

#: Targets a link's text repeats, by a made-up label.  markdownify wrote each
#: as `<target>`, which Python-Markdown links only for http, https and ftp; it
#: read the rest as an HTML tag, which nh3 dropped with the text, or kept and
#: let run on.  The https target is the control, whose page did not change.
#: The tag names come last: `script` and the rest took what followed them.
BARE_LINK_TARGETS = {
    "pdf": "quillmoss.pdf",
    "xhtml": "chapter-vornet.xhtml",
    "folder": "notes/brindlow.txt",
    "root": "/sorvane/index.html",
    "host": "//example.invalid/fenwold",
    "query": "?glimt=1",
    "www": "www.example.invalid",
    "file": "file:///tmp/tarnby.pdf",
    "https": "https://example.invalid/marrowby",
    "em": "em",
    "br": "br",
    "strong": "strong",
    "a": "a",
    "img": "img",
    "li": "li",
    "pre": "pre",
    "table": "table",
    "title": "title",
    "textarea": "textarea",
    "style": "style",
    "script": "script",
}
#: Where a book or a page puts a link besides a paragraph, and which target it
#: links there.  A link in code is written as its text: a code span holds none.
_BARE_LINK_PLACES = (
    ("heading", "pdf"),
    ("item", "pdf"),
    ("cell", "pdf"),
    ("code", "pdf"),
    ("code", "https"),
    ("kbd", "em"),
    ("samp", "folder"),
)
#: Word has no heading, list or code here without a style map: a table cell.
_WORD_BARE_LINK_PLACES = (("cell", "pdf"),)
_BARE_LINK_HTML = {
    "para": "<p>{0} {1} {2}</p>",
    "heading": "<h2>{0} {1} {2}</h2>",
    "item": "<ul><li>{0} {1} {2}</li></ul>",
    "cell": "<table><tr><td>{0} {1} {2}</td></tr></table>",
    "code": "<p>{0} <code>{1}</code> {2}</p>",
    "kbd": "<p>{0} <kbd>{1}</kbd> {2}</p>",
    "samp": "<p>{0} <samp>{1}</samp> {2}</p>",
}


@dataclass(frozen=True)
class BareLinkLine:
    """One block of a `bare_link_*` fixture: a made-up word, a link whose text
    is its target, and another made-up word, in the block `where` names."""

    label: str
    target: str
    where: str

    @property
    def words(self) -> tuple[str, str, str]:
        stem = f"{self.label}{self.where}"
        return (f"{stem}prior", self.target, f"{stem}next")

    @property
    def text(self) -> str:
        """The block as the page must read it."""
        return " ".join(self.words)


def bare_link_lines(name: str) -> list[BareLinkLine]:
    """The blocks of the `bare_link_*` fixture called `name`: each place, then
    a paragraph for every target."""
    places = _WORD_BARE_LINK_PLACES if name.endswith(".docx") else _BARE_LINK_PLACES
    return [
        BareLinkLine(label, BARE_LINK_TARGETS[label], where) for where, label in places
    ] + [
        BareLinkLine(label, target, "para")
        for label, target in BARE_LINK_TARGETS.items()
    ]


def write_bare_link_report(path: Path) -> None:
    lines = bare_link_lines("bare_link_report.docx")
    ids = {
        target: f"rId{index}"
        for index, target in enumerate(dict.fromkeys(line.target for line in lines), 2)
    }
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(
            f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"'
            f' Target="{target}" TargetMode="External"/>'
            for target, rid in ids.items()
        )
        + "</Relationships>\n"
    )

    def block(line: BareLinkLine) -> str:
        prior, target, following = line.words
        link = f'<w:hyperlink r:id="{ids[target]}">{_word_run(target)}</w:hyperlink>'
        paragraph = (
            f"<w:p>{_word_run(prior + ' ')}{link}{_word_run(' ' + following)}</w:p>"
        )
        if line.where == "cell":
            return f"<w:tbl><w:tblPr/><w:tr><w:tc><w:tcPr/>{paragraph}</w:tc></w:tr></w:tbl>"
        return paragraph

    body = _docx_paragraph(SENTINELS["bare_link_report.docx"]) + "".join(
        map(block, lines)
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {_DECLARE}><w:body>{body}</w:body></w:document>\n"
    )
    _write_archive(
        path,
        [
            ("[Content_Types].xml", _DOCX_CONTENT_TYPES),
            ("_rels/.rels", _DOCX_RELS),
            ("word/document.xml", document),
            ("word/_rels/document.xml.rels", relationships),
        ],
    )


def _bare_link_html(name: str) -> str:
    return "\n".join(
        "  "
        + _BARE_LINK_HTML[line.where].format(
            line.words[0],
            f'<a href="{line.target}">{line.target}</a>',
            line.words[2],
        )
        for line in bare_link_lines(name)
    )


def write_bare_link_book(path: Path) -> None:
    opf = _EPUB_OPF.replace("webshot-office-fixture", "webshot-bare-link-fixture")
    chapter = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter One</title></head>
<body>
  <p>{SENTINELS["bare_link_book.epub"]}</p>
{_bare_link_html("bare_link_book.epub")}
</body></html>
"""
    _write_archive(
        path,
        [
            ("mimetype", "application/epub+zip"),
            ("META-INF/container.xml", _EPUB_CONTAINER),
            ("OEBPS/content.opf", opf),
            ("OEBPS/chapter.xhtml", chapter),
        ],
        first_stored=True,
    )


def write_bare_link_archive(path: Path) -> None:
    """An HTML page in a zip, as `write_edge_space_archive` explains."""
    page = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Bare links</title></head>
<body>
  <p>{SENTINELS["bare_link_archive.zip"]}</p>
{_bare_link_html("bare_link_archive.zip")}
</body></html>
"""
    _write_archive(path, [("pages/bare-links.html", page)])


WRITERS = {
    "sample_report.docx": write_docx,
    "sample_deck.pptx": write_pptx,
    "sample_workbook.xlsx": write_xlsx,
    "sample_book.epub": write_epub,
    "numbered_book.epub": write_numbered_epub,
    "sample_archive.zip": write_zip,
    "equation_deck.pptx": write_equation_deck,
    "equation_report.docx": write_equation_report,
    "equation_archive.zip": write_equation_archive,
    "markup_deck.pptx": write_markup_deck,
    "markup_report.docx": write_markup_report,
    "markup_workbook.xlsx": write_markup_workbook,
    "markup_book.epub": write_markup_book,
    "markup_archive.zip": write_markup_archive,
    "left_out_archive.zip": write_left_out_archive,
    "adjacent_lists_book.epub": write_adjacent_lists_book,
    "picture_deck.pptx": write_picture_deck,
    "picture_report.docx": write_picture_report,
    "picture_book.epub": write_picture_book,
    "late_utf8.csv": write_late_utf8_csv,
    "late_utf8_archive.zip": write_late_utf8_archive,
    "nested_book.epub": write_nested_book,
    "nested_report.docx": write_nested_report,
    "nested_archive.zip": write_nested_archive,
    "code_items_book.epub": write_code_items_book,
    "code_items_report.docx": write_code_items_report,
    "code_items_archive.zip": write_code_items_archive,
    "definition_book.epub": write_definition_book,
    "definition_archive.zip": write_definition_archive,
    "inline_code_book.epub": write_inline_code_book,
    "inline_code_report.docx": write_inline_code_report,
    "inline_code_archive.zip": write_inline_code_archive,
    "quoted_code_book.epub": write_quoted_code_book,
    "quoted_code_report.docx": write_quoted_code_report,
    "quoted_code_archive.zip": write_quoted_code_archive,
    "edge_space_report.docx": write_edge_space_report,
    "edge_space_book.epub": write_edge_space_book,
    "edge_space_archive.zip": write_edge_space_archive,
    "whitespace_run_report.docx": write_whitespace_run_report,
    "whitespace_run_book.epub": write_whitespace_run_book,
    "whitespace_run_archive.zip": write_whitespace_run_archive,
    "adjacent_blocks_book.epub": write_adjacent_blocks_book,
    "adjacent_blocks_report.docx": write_adjacent_blocks_report,
    "adjacent_blocks_archive.zip": write_adjacent_blocks_archive,
    "numbered_report.docx": write_numbered_report,
    "definitions_book.epub": write_definitions_book,
    "definitions_archive.zip": write_definitions_archive,
    "quoted_lists_book.epub": write_quoted_lists_book,
    "quoted_lists_archive.zip": write_quoted_lists_archive,
    "line_break_report.docx": write_line_break_report,
    "line_break_book.epub": write_line_break_book,
    "line_break_archive.zip": write_line_break_archive,
    "bare_link_report.docx": write_bare_link_report,
    "bare_link_book.epub": write_bare_link_book,
    "bare_link_archive.zip": write_bare_link_archive,
}


def main() -> int:
    for name, writer in WRITERS.items():
        destination = FIXTURE_ROOT / name
        writer(destination)
        print(f"wrote {destination.relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
